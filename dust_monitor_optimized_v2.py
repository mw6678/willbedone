import os
import sys
import time
import sqlite3
import traceback
import queue
import threading
import json
import math
import serial
import serial.tools.list_ports

from collections import deque
from datetime import datetime, timedelta

from PyQt5.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QLabel, QFrame, QPushButton, QDialog, QMessageBox, QScrollArea,
    QGroupBox, QComboBox, QLineEdit
)
from PyQt5.QtCore import QThread, pyqtSignal, Qt, QPoint
from PyQt5.QtGui import QFont

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


BASE_DIR = os.path.dirname(os.path.abspath(__file__))


# ============================================================
# Config
# ============================================================

class Config:
    BAUD_RATE = 9600
    SERIAL_TIMEOUT = 0.1

    MAX_SENSORS = 4

    # UI 평균용
    SMOOTHING_WINDOW = 1

    # 일정 시간 데이터가 없을 경우 오류
    NO_DATA_TIMEOUT = 30.0

    # 연결 끊김 후 재접속
    RECONNECT_DELAY_MS = 3000

    # 오래된 로그 삭제
    LOG_RETENTION_DAYS = 30

    # Thread 종료 대기
    THREAD_WAIT_MS = 5000

    # DB
    DB_QUEUE_SIZE = 2000

    # 디렉터리
    DATA_DIR = os.path.join(BASE_DIR, "Data")
    LOG_DIR = os.path.join(BASE_DIR, "Logs")
    EXCEL_DIR = os.path.join(BASE_DIR, "Excel_Logs")

    CALIBRATION_FILE = os.path.join(
        DATA_DIR,
        "sensor_calibration.json"
    )

    VACUUM_MIN_INTERVAL_DAYS = 7

    # --------------------------------------------------------
    # 기본 보정값
    #
    # scale:
    #   PM10, PM2.5, PM1.0 순서
    #
    # offset:
    #   PM10, PM2.5, PM1.0 순서
    #
    # 최종값:
    #   PM10  = raw_pm10  * scale[0] + offset[0]
    #   PM2.5 = raw_pm25  * scale[1] + offset[1]
    #   PM1.0 = raw_pm1   * scale[2] + offset[2]
    # --------------------------------------------------------

    DEFAULT_SENSOR_CALIBRATION = {
        0: {
            "scale": (1.0, 2.0, 1.0),
            "offset": (0.0, 1.0, 0.0),
        },
        1: {
            "scale": (1.0, 3.0, 1.0),
            "offset": (0.0, 1.0, 0.0),
        },
        2: {
            "scale": (1.0, 1.0, 1.0),
            "offset": (0.0, 0.0, 0.0),
        },
        3: {
            "scale": (1.0, 1.0, 1.0),
            "offset": (0.0, 0.0, 0.0),
        },
    }

    SENSOR_CALIBRATION = {}

    # --------------------------------------------------------
    # 미세먼지 등급
    # --------------------------------------------------------

    PM10_LEVELS = [
        {"name": "좋음", "min": 0, "max": 30, "color": "#28A745"},
        {"name": "보통", "min": 31, "max": 80, "color": "#FFD700"},
        {"name": "민감군", "min": 81, "max": 120, "color": "#FD7E14"},
        {"name": "나쁨", "min": 121, "max": 150, "color": "#DC3545"},
        {"name": "매우 나쁨", "min": 151, "max": 300, "color": "#800080"},
        {"name": "위험", "min": 301, "max": 600, "color": "#795548"},
    ]

    PM25_LEVELS = [
        {"name": "좋음", "min": 0, "max": 15, "color": "#28A745"},
        {"name": "보통", "min": 16, "max": 35, "color": "#FFD700"},
        {"name": "민감군", "min": 36, "max": 50, "color": "#FD7E14"},
        {"name": "나쁨", "min": 51, "max": 75, "color": "#DC3545"},
        {"name": "매우 나쁨", "min": 76, "max": 100, "color": "#800080"},
        {"name": "위험", "min": 101, "max": 500, "color": "#795548"},
    ]

    PM1_LEVELS = [
        {"name": "좋음", "min": 0, "max": 10, "color": "#28A745"},
        {"name": "보통", "min": 11, "max": 25, "color": "#FFD700"},
        {"name": "민감군", "min": 26, "max": 35, "color": "#FD7E14"},
        {"name": "나쁨", "min": 36, "max": 50, "color": "#DC3545"},
        {"name": "매우 나쁨", "min": 51, "max": 75, "color": "#800080"},
        {"name": "위험", "min": 76, "max": 300, "color": "#795548"},
    ]

    @classmethod
    def _normalize_calibration(cls, data):
        normalized = {}

        defaults = cls.DEFAULT_SENSOR_CALIBRATION

        for i in range(cls.MAX_SENSORS):

            if isinstance(data, dict):
                raw = data.get(
                    str(i),
                    data.get(i, defaults.get(i))
                )
            else:
                raw = defaults.get(i)

            try:
                scales = tuple(
                    float(x) for x in raw["scale"][:3]
                )

                offsets = tuple(
                    float(x) for x in raw["offset"][:3]
                )

                if len(scales) != 3 or len(offsets) != 3:
                    raise ValueError

                # NaN / Inf 방지
                if not all(math.isfinite(x) for x in scales):
                    raise ValueError

                if not all(math.isfinite(x) for x in offsets):
                    raise ValueError

                # Scale은 음수 허용하지 않음
                if any(x < 0 for x in scales):
                    raise ValueError

                normalized[i] = {
                    "scale": scales,
                    "offset": offsets,
                }

            except (
                KeyError,
                TypeError,
                ValueError,
                IndexError
            ):
                normalized[i] = {
                    "scale": tuple(defaults[i]["scale"]),
                    "offset": tuple(defaults[i]["offset"]),
                }

        return normalized

    @classmethod
    def load_calibration(cls):
        os.makedirs(cls.DATA_DIR, exist_ok=True)

        try:
            if os.path.exists(cls.CALIBRATION_FILE):

                with open(
                    cls.CALIBRATION_FILE,
                    "r",
                    encoding="utf-8"
                ) as f:
                    data = json.load(f)

                cls.SENSOR_CALIBRATION = cls._normalize_calibration(data)
                return

        except (
            OSError,
            json.JSONDecodeError
        ) as e:
            print(f"보정값 파일 읽기 오류: {e}")

        cls.SENSOR_CALIBRATION = cls._normalize_calibration(
            cls.DEFAULT_SENSOR_CALIBRATION
        )

        cls.save_calibration()

    @classmethod
    def save_calibration(cls):

        os.makedirs(cls.DATA_DIR, exist_ok=True)

        serializable = {
            str(i): {
                "scale": list(params["scale"]),
                "offset": list(params["offset"]),
            }
            for i, params in cls.SENSOR_CALIBRATION.items()
        }

        temp_file = cls.CALIBRATION_FILE + ".tmp"

        try:
            with open(
                temp_file,
                "w",
                encoding="utf-8"
            ) as f:

                json.dump(
                    serializable,
                    f,
                    ensure_ascii=False,
                    indent=2
                )

            os.replace(
                temp_file,
                cls.CALIBRATION_FILE
            )

            return True

        except OSError as e:

            print(f"보정값 파일 저장 오류: {e}")

            try:
                if os.path.exists(temp_file):
                    os.remove(temp_file)
            except OSError:
                pass

            return False


# ============================================================
# Dust Parser
# ============================================================

class DustParser:

    @staticmethod
    def parse(raw_data, calib_params):

        try:

            raw_data = raw_data.strip()

            if not raw_data:
                return None

            parts = [
                item.strip()
                for item in raw_data.split(",")
            ]

            if len(parts) < 3:
                return None

            scales = calib_params.get(
                "scale",
                (1.0, 1.0, 1.0)
            )

            offsets = calib_params.get(
                "offset",
                (0.0, 0.0, 0.0)
            )

            # ------------------------------------------------
            # 센서 Raw 데이터
            #
            # parts[0] = PM1.0
            # parts[1] = PM2.5
            # parts[2] = PM10
            # ------------------------------------------------

            raw_pm1 = float(parts[0])
            raw_pm25 = float(parts[1])
            raw_pm10 = float(parts[2])

            # ------------------------------------------------
            # 보정값
            # ------------------------------------------------

            scale_pm10 = float(scales[0])
            scale_pm25 = float(scales[1])
            scale_pm1 = float(scales[2])

            offset_pm10 = float(offsets[0])
            offset_pm25 = float(offsets[1])
            offset_pm1 = float(offsets[2])

            values = (
                raw_pm1,
                raw_pm25,
                raw_pm10,
                scale_pm10,
                scale_pm25,
                scale_pm1,
                offset_pm10,
                offset_pm25,
                offset_pm1,
            )

            if not all(
                math.isfinite(v)
                for v in values
            ):
                return None

            # =================================================
            # 핵심
            #
            # 반드시
            #
            # 최종값 = 원본값 × Scale + Offset
            #
            # 순서대로 계산한다.
            # =================================================

            calibrated_pm10 = (
                raw_pm10 * scale_pm10
            ) + offset_pm10

            calibrated_pm25 = (
                raw_pm25 * scale_pm25
            ) + offset_pm25

            calibrated_pm1 = (
                raw_pm1 * scale_pm1
            ) + offset_pm1

            # 음수 방지
            pm10 = max(
                0,
                int(round(calibrated_pm10))
            )

            pm25 = max(
                0,
                int(round(calibrated_pm25))
            )

            pm1 = max(
                0,
                int(round(calibrated_pm1))
            )

            # 비정상 범위
            if (
                pm1 > 1000
                or pm25 > 1000
                or pm10 > 2000
            ):
                return "OUT_OF_RANGE"

            return {
                "pm10": pm10,
                "pm25": pm25,
                "pm1": pm1,
            }

        except (
            ValueError,
            TypeError,
            IndexError
        ):
            return None


# ============================================================
# Database Manager
# ============================================================

class DatabaseManager:

    _instance = None
    _lock = threading.Lock()

    def __new__(
        cls,
        db_path=None
    ):

        with cls._lock:

            if cls._instance is None:

                cls._instance = super(
                    DatabaseManager,
                    cls
                ).__new__(cls)

                cls._instance._initialized = False

        return cls._instance

    def __init__(
        self,
        db_path=None
    ):

        if self._initialized:
            return

        with self._lock:

            if self._initialized:
                return

            self.db_path = (
                db_path
                or os.path.join(
                    BASE_DIR,
                    "Data",
                    "dust_measurement.db"
                )
            )

            self.queue = queue.Queue(
                maxsize=Config.DB_QUEUE_SIZE
            )

            self.running = True

            self._stop_lock = threading.Lock()
            self._stopped = False

            self._init_db()

            self.worker_thread = threading.Thread(
                target=self._db_worker,
                name="SQLiteWorker",
                daemon=True
            )

            self.worker_thread.start()

            self._initialized = True

    def _connect(self):

        conn = sqlite3.connect(
            self.db_path,
            timeout=15.0
        )

        conn.execute(
            "PRAGMA busy_timeout=10000;"
        )

        return conn

    def _init_db(self):

        try:

            os.makedirs(
                os.path.dirname(
                    self.db_path
                ),
                exist_ok=True
            )

            with sqlite3.connect(
                self.db_path,
                timeout=15.0
            ) as conn:

                conn.execute(
                    "PRAGMA journal_mode=WAL;"
                )

                conn.execute(
                    "PRAGMA synchronous=NORMAL;"
                )

                conn.execute("""
                    CREATE TABLE IF NOT EXISTS measurements (
                        id INTEGER PRIMARY KEY AUTOINCREMENT,
                        measured_at TEXT NOT NULL,
                        sensor_index INTEGER NOT NULL,
                        port TEXT NOT NULL,
                        pm10 REAL,
                        pm25 REAL,
                        pm1 REAL,
                        status TEXT DEFAULT 'NORMAL'
                    )
                """)

                conn.execute("""
                    CREATE INDEX IF NOT EXISTS
                    idx_measurements_time
                    ON measurements(measured_at)
                """)

                conn.execute("""
                    CREATE INDEX IF NOT EXISTS
                    idx_measurements_sensor
                    ON measurements(sensor_index)
                """)

                conn.execute("""
                    CREATE INDEX IF NOT EXISTS
                    idx_measurements_sensor_time
                    ON measurements(
                        sensor_index,
                        measured_at
                    )
                """)

                conn.commit()

        except Exception as e:

            print(
                f"Global SQLite 초기화 오류: {e}"
            )

    def _db_worker(self):

        conn = None

        try:

            conn = self._connect()

            while True:

                try:

                    task = self.queue.get(
                        timeout=0.5
                    )

                except queue.Empty:

                    if (
                        not self.running
                        and self.queue.empty()
                    ):
                        break

                    continue

                if task is None:

                    self.queue.task_done()

                    if (
                        not self.running
                        and self.queue.empty()
                    ):
                        break

                    continue

                (
                    func,
                    args,
                    future_event,
                    result_holder
                ) = task

                try:

                    result = func(
                        conn,
                        *args
                    )

                    conn.commit()

                    if result_holder is not None:
                        result_holder["result"] = result

                except Exception as e:

                    try:
                        conn.rollback()
                    except Exception:
                        pass

                    print(
                        f"DB Worker 작업 수행 오류: {e}"
                    )

                    if result_holder is not None:
                        result_holder["error"] = e

                finally:

                    if future_event is not None:
                        future_event.set()

                    self.queue.task_done()

        except Exception as e:

            print(
                f"DB Worker 초기화/루프 오류: {e}"
            )

        finally:

            if conn is not None:

                try:
                    conn.close()
                except Exception:
                    pass

    def execute_async(
        self,
        func,
        *args
    ):

        with self._stop_lock:

            if self._stopped:
                return False

            try:

                self.queue.put_nowait(
                    (
                        func,
                        args,
                        None,
                        None
                    )
                )

                return True

            except queue.Full:

                print(
                    "DB 작업 큐가 가득 찼습니다."
                )

                return False

    def execute_sync(
        self,
        func,
        *args,
        timeout=5.0
    ):

        future_event = threading.Event()
        result_holder = {}

        with self._stop_lock:

            if self._stopped:
                return None

            try:

                self.queue.put(
                    (
                        func,
                        args,
                        future_event,
                        result_holder
                    ),
                    timeout=2.0
                )

            except queue.Full:

                print(
                    "DB 작업 큐가 가득 찼습니다."
                )

                return None

        if not future_event.wait(
            timeout=timeout
        ):

            print(
                "DB 동기 작업 시간 초과"
            )

            return None

        if "error" in result_holder:
            return None

        return result_holder.get(
            "result"
        )

    def flush(
        self,
        timeout=10.0
    ):

        done = threading.Event()
        result_holder = {}

        def barrier(conn):
            return True

        with self._stop_lock:

            if self._stopped:
                return False

            try:

                self.queue.put(
                    (
                        barrier,
                        (),
                        done,
                        result_holder
                    ),
                    timeout=2.0
                )

            except queue.Full:

                return False

        if not done.wait(
            timeout=timeout
        ):
            return False

        return "error" not in result_holder

    def stop(
        self,
        timeout=10.0
    ):

        with self._stop_lock:

            if self._stopped:
                return True

            self.running = False

        # 큐의 작업을 먼저 처리하고 종료
        try:
            self.queue.put(
                None,
                timeout=2.0
            )
        except queue.Full:
            pass

        self.worker_thread.join(
            timeout=timeout
        )

        if self.worker_thread.is_alive():

            print(
                "DB Worker 종료 시간 초과"
            )

            return False

        with self._stop_lock:
            self._stopped = True

        return True


# ============================================================
# Sensor Logger
# ============================================================

class SensorLogger:

    _vacuum_lock = threading.Lock()

    def __init__(
        self,
        port_name,
        sensor_index
    ):

        self.port_name = port_name
        self.sensor_index = sensor_index

        self.log_dir = Config.LOG_DIR
        self.excel_dir = Config.EXCEL_DIR
        self.db_dir = Config.DATA_DIR

        os.makedirs(
            self.log_dir,
            exist_ok=True
        )

        os.makedirs(
            self.excel_dir,
            exist_ok=True
        )

        os.makedirs(
            self.db_dir,
            exist_ok=True
        )

        self.db_path = os.path.join(
            self.db_dir,
            "dust_measurement.db"
        )

        self.db_manager = DatabaseManager(
            self.db_path
        )

    @staticmethod
    def init_global_database(
        db_path
    ):

        DatabaseManager(db_path)

    @staticmethod
    def cleanup_old_logs_global():

        threshold_time = (
            time.time()
            - (
                Config.LOG_RETENTION_DAYS
                * 24
                * 60
                * 60
            )
        )

        target_dirs = [
            Config.LOG_DIR,
            Config.EXCEL_DIR,
        ]

        for target_dir in target_dirs:

            if not os.path.exists(
                target_dir
            ):
                continue

            try:
                filenames = os.listdir(
                    target_dir
                )
            except OSError:
                continue

            for filename in filenames:

                file_path = os.path.join(
                    target_dir,
                    filename
                )

                if not os.path.isfile(
                    file_path
                ):
                    continue

                try:

                    if (
                        os.path.getmtime(
                            file_path
                        )
                        < threshold_time
                    ):
                        os.remove(file_path)

                except OSError:
                    pass

        # 오래된 DB 데이터 삭제
        db_path = os.path.join(
            Config.DATA_DIR,
            "dust_measurement.db"
        )

        if not os.path.exists(
            db_path
        ):
            return

        try:

            threshold_datetime = (
                datetime.fromtimestamp(
                    threshold_time
                ).strftime(
                    "%Y-%m-%d %H:%M:%S"
                )
            )

            def delete_old(conn):

                cursor = conn.execute(
                    """
                    DELETE FROM measurements
                    WHERE measured_at < ?
                    """,
                    (threshold_datetime,)
                )

                return cursor.rowcount

            db_mgr = DatabaseManager(
                db_path
            )

            deleted_count = db_mgr.execute_sync(
                delete_old,
                timeout=15.0
            )

            # 삭제량이 많을 때만 VACUUM
            if (
                isinstance(
                    deleted_count,
                    int
                )
                and deleted_count > 1000
            ):

                vacuum_marker = os.path.join(
                    Config.DATA_DIR,
                    ".last_vacuum"
                )

                do_vacuum = True

                try:

                    last_vacuum = (
                        os.path.getmtime(
                            vacuum_marker
                        )
                    )

                    do_vacuum = (
                        time.time()
                        - last_vacuum
                        >=
                        Config.VACUUM_MIN_INTERVAL_DAYS
                        * 86400
                    )

                except OSError:
                    pass

                if do_vacuum:

                    def vacuum(conn):
                        conn.execute(
                            "VACUUM"
                        )
                        return True

                    result = db_mgr.execute_sync(
                        vacuum,
                        timeout=60.0
                    )

                    if result:

                        try:

                            with open(
                                vacuum_marker,
                                "a"
                            ):
                                pass

                        except OSError:
                            pass

        except Exception as e:

            print(
                f"DB 오래된 데이터 삭제 오류: {e}"
            )

    def write_error(
        self,
        message
    ):

        log_file = os.path.join(
            self.log_dir,
            f"error_log_Sensor"
            f"{self.sensor_index + 1}.txt"
        )

        timestamp = datetime.now().strftime(
            "%Y-%m-%d %H:%M:%S"
        )

        try:

            with open(
                log_file,
                "a",
                encoding="utf-8-sig"
            ) as f:

                f.write(
                    f"{timestamp} | "
                    f"포트: {self.port_name} | "
                    f"{message}\n"
                )

        except OSError:
            pass

    def save_measurements_batch(
        self,
        measurements
    ):

        if not measurements:
            return True

        data = list(measurements)

        def insert_rows(
            conn,
            rows
        ):

            conn.executemany(
                """
                INSERT INTO measurements (
                    measured_at,
                    sensor_index,
                    port,
                    pm10,
                    pm25,
                    pm1,
                    status
                )
                VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                rows
            )

        return self.db_manager.execute_async(
            insert_rows,
            data
        )

    def save_measurement_sync(
        self,
        measurement,
        timeout=10.0
    ):

        def insert_one(
            conn,
            row
        ):

            conn.execute(
                """
                INSERT INTO measurements (
                    measured_at,
                    sensor_index,
                    port,
                    pm10,
                    pm25,
                    pm1,
                    status
                )
                VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                row
            )

            return True

        return bool(
            self.db_manager.execute_sync(
                insert_one,
                measurement,
                timeout=timeout
            )
        )

    def _fetch_rows_for_date(
        self,
        date_str
    ):

        start_datetime = (
            f"{date_str} 00:00:00"
        )

        next_date = (
            datetime.strptime(
                date_str,
                "%Y-%m-%d"
            )
            + timedelta(days=1)
        ).strftime(
            "%Y-%m-%d"
        )

        end_datetime = (
            f"{next_date} 00:00:00"
        )

        def fetch_rows(conn):

            cursor = conn.execute(
                """
                SELECT
                    measured_at,
                    port,
                    pm10,
                    pm25,
                    pm1,
                    status
                FROM measurements
                WHERE
                    sensor_index = ?
                    AND measured_at >= ?
                    AND measured_at < ?
                ORDER BY measured_at
                """,
                (
                    self.sensor_index,
                    start_datetime,
                    end_datetime
                )
            )

            return cursor.fetchall()

        return self.db_manager.execute_sync(
            fetch_rows,
            timeout=15.0
        )

    def _write_excel(
        self,
        date_str,
        rows
    ):

        if not rows:
            return False

        filename = (
            f"Dust_log_{date_str}"
            f"_Sensor{self.sensor_index + 1}.xlsx"
        )

        file_path = os.path.join(
            self.excel_dir,
            filename
        )

        temp_path = (
            file_path + ".tmp"
        )

        wb = Workbook()

        ws = wb.active

        ws.title = (
            f"{date_str} 측정 데이터"
        )

        headers = [
            "측정일시",
            "포트",
            "PM10",
            "PM2.5",
            "PM1.0",
            "상태"
        ]

        ws.append(headers)

        # 스타일
        thin_border = Border(
            left=Side(
                style="thin",
                color="D3D3D3"
            ),
            right=Side(
                style="thin",
                color="D3D3D3"
            ),
            top=Side(
                style="thin",
                color="D3D3D3"
            ),
            bottom=Side(
                style="thin",
                color="D3D3D3"
            ),
        )

        header_fill = PatternFill(
            start_color="1F4E78",
            end_color="1F4E78",
            fill_type="solid"
        )

        header_font = Font(
            name="Malgun Gothic",
            size=10,
            bold=True,
            color="FFFFFF"
        )

        data_font = Font(
            name="Malgun Gothic",
            size=9
        )

        center_align = Alignment(
            horizontal="center",
            vertical="center"
        )

        right_align = Alignment(
            horizontal="right",
            vertical="center"
        )

        # 헤더 스타일
        for col_num in range(
            1,
            len(headers) + 1
        ):

            cell = ws.cell(
                row=1,
                column=col_num
            )

            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border

        # 데이터
        for row_data in rows:

            ws.append(
                list(row_data)
            )

            current_row = ws.max_row

            for col_num, value in enumerate(
                row_data,
                1
            ):

                cell = ws.cell(
                    row=current_row,
                    column=col_num
                )

                cell.font = data_font
                cell.border = thin_border

                if col_num in (
                    1,
                    2,
                    6
                ):

                    cell.alignment = center_align

                else:

                    cell.alignment = right_align

                    if isinstance(
                        value,
                        (int, float)
                    ):

                        cell.number_format = "#,##0"

        widths = [
            19,
            14,
            10,
            10,
            10,
            12
        ]

        for i, width in enumerate(
            widths,
            1
        ):

            ws.column_dimensions[
                get_column_letter(i)
            ].width = width

        try:

            wb.save(
                temp_path
            )

            os.replace(
                temp_path,
                file_path
            )

            return True

        finally:

            try:

                if os.path.exists(
                    temp_path
                ):
                    os.remove(
                        temp_path
                    )

            except OSError:
                pass

    def export_excel(
        self,
        date_str=None
    ):

        try:

            date_str = (
                date_str
                or datetime.now().strftime(
                    "%Y-%m-%d"
                )
            )

            rows = self._fetch_rows_for_date(
                date_str
            )

            if not rows:
                return False

            return self._write_excel(
                date_str,
                rows
            )

        except Exception as e:

            self.write_error(
                f"Excel Export 오류: {e}"
            )

            return False

    def export_all_dates_excel(self):

        try:

            def get_dates(conn):

                cursor = conn.execute(
                    """
                    SELECT DISTINCT
                        SUBSTR(measured_at, 1, 10)
                    FROM measurements
                    WHERE sensor_index = ?
                    ORDER BY measured_at
                    """,
                    (self.sensor_index,)
                )

                return [
                    row[0]
                    for row in cursor.fetchall()
                ]

            dates = self.db_manager.execute_sync(
                get_dates,
                timeout=15.0
            )

            if not dates:
                return False

            success_any = False

            for date_str in dates:

                if self.export_excel(
                    date_str
                ):
                    success_any = True

            return success_any

        except Exception as e:

            self.write_error(
                f"전체 날짜 엑셀 내보내기 오류: {e}"
            )

            return False


# ============================================================
# Serial Thread
# ============================================================

class SerialThread(QThread):

    data_signal = pyqtSignal(
        int,
        int,
        int,
        int
    )

    error_signal = pyqtSignal(
        int,
        str
    )

    # 날짜 변경 알림
    #
    # sensor_index
    # previous_date
    #
    # 예:
    # 2026-09-09 -> 2026-09-10
    #
    # 2026-09-09 Excel 자동 저장
    date_changed_signal = pyqtSignal(
        int,
        str
    )

    def __init__(
        self,
        port_name,
        sensor_index
    ):

        super().__init__()

        self.port_name = port_name
        self.sensor_index = sensor_index

        default_calib = {
            "scale": (
                1.0,
                1.0,
                1.0
            ),
            "offset": (
                0.0,
                0.0,
                0.0
            ),
        }

        self.calib_params = self._copy_calibration(
            Config.SENSOR_CALIBRATION.get(
                sensor_index,
                default_calib
            )
        )

        self.logger = SensorLogger(
            port_name,
            sensor_index
        )

        window_size = max(
            1,
            Config.SMOOTHING_WINDOW
        )

        self.data_buffer = {
            "pm10": deque(
                maxlen=window_size
            ),
            "pm25": deque(
                maxlen=window_size
            ),
            "pm1": deque(
                maxlen=window_size
            ),
        }

        self.minute_sum = {
            "pm10": 0,
            "pm25": 0,
            "pm1": 0
        }

        self.minute_count = 0

        self.current_minute_bucket = (
            self._get_current_minute()
        )

        self._buffer_lock = threading.Lock()

    @staticmethod
    def _copy_calibration(
        calib
    ):

        return {
            "scale": tuple(
                float(x)
                for x in calib["scale"]
            ),
            "offset": tuple(
                float(x)
                for x in calib["offset"]
            )
        }

    @staticmethod
    def _get_current_minute():

        return datetime.now().replace(
            second=0,
            microsecond=0
        )

    def safe_sleep(
        self,
        milliseconds
    ):

        remaining = milliseconds

        while remaining > 0:

            if self.isInterruptionRequested():
                return

            sleep_time = min(
                50,
                remaining
            )

            self.msleep(
                sleep_time
            )

            remaining -= sleep_time

    # --------------------------------------------------------
    # Buffer
    # --------------------------------------------------------

    def clear_ui_buffers(self):

        with self._buffer_lock:

            for buffer in self.data_buffer.values():
                buffer.clear()

    def clear_minute_buffers(self):

        with self._buffer_lock:

            self.minute_sum["pm10"] = 0
            self.minute_sum["pm25"] = 0
            self.minute_sum["pm1"] = 0

            self.minute_count = 0

    def clear_all_buffers(self):

        with self._buffer_lock:

            for buffer in self.data_buffer.values():
                buffer.clear()

            self.minute_sum["pm10"] = 0
            self.minute_sum["pm25"] = 0
            self.minute_sum["pm1"] = 0

            self.minute_count = 0

    def add_measurement(
        self,
        pm10,
        pm25,
        pm1
    ):

        with self._buffer_lock:

            self.data_buffer[
                "pm10"
            ].append(pm10)

            self.data_buffer[
                "pm25"
            ].append(pm25)

            self.data_buffer[
                "pm1"
            ].append(pm1)

            self.minute_sum[
                "pm10"
            ] += pm10

            self.minute_sum[
                "pm25"
            ] += pm25

            self.minute_sum[
                "pm1"
            ] += pm1

            self.minute_count += 1

    # --------------------------------------------------------
    # 현재 분 평균 저장
    # --------------------------------------------------------

    def save_current_minute_average(
        self,
        minute_bucket,
        synchronous=False
    ):

        with self._buffer_lock:

            count = self.minute_count

            if count <= 0:
                return True

            avg_pm10 = int(
                round(
                    self.minute_sum["pm10"]
                    / count
                )
            )

            avg_pm25 = int(
                round(
                    self.minute_sum["pm25"]
                    / count
                )
            )

            avg_pm1 = int(
                round(
                    self.minute_sum["pm1"]
                    / count
                )
            )

        measured_at = (
            minute_bucket.strftime(
                "%Y-%m-%d %H:%M:00"
            )
        )

        row = (
            measured_at,
            self.sensor_index,
            self.port_name,
            avg_pm10,
            avg_pm25,
            avg_pm1,
            "NORMAL",
        )

        if synchronous:

            success = (
                self.logger.save_measurement_sync(
                    row,
                    timeout=10.0
                )
            )

        else:

            success = (
                self.logger.save_measurements_batch(
                    [row]
                )
            )

        if success:

            with self._buffer_lock:

                # 저장 중 새로운 데이터가 들어왔다면
                # 새 데이터는 삭제하지 않는다.
                if self.minute_count == count:

                    self.minute_sum[
                        "pm10"
                    ] = 0

                    self.minute_sum[
                        "pm25"
                    ] = 0

                    self.minute_sum[
                        "pm1"
                    ] = 0

                    self.minute_count = 0

            return True

        return False

    # --------------------------------------------------------
    # 분 / 날짜 경계 처리
    # --------------------------------------------------------

    def check_minute_boundary(self):

        new_bucket = (
            self._get_current_minute()
        )

        if (
            new_bucket
            == self.current_minute_bucket
        ):
            return

        old_bucket = (
            self.current_minute_bucket
        )

        # 이전 분 저장
        if self.minute_count > 0:

            saved = (
                self.save_current_minute_average(
                    old_bucket,
                    synchronous=False
                )
            )

            if not saved:
                return

        # ----------------------------------------------------
        # 날짜가 바뀌었는지 확인
        # ----------------------------------------------------

        if (
            old_bucket.date()
            != new_bucket.date()
        ):

            previous_date = (
                old_bucket.strftime(
                    "%Y-%m-%d"
                )
            )

            # 앱에 날짜 변경 알림
            #
            # 앱에서 이전 날짜 Excel을 자동 생성한다.
            self.date_changed_signal.emit(
                self.sensor_index,
                previous_date
            )

        self.current_minute_bucket = new_bucket

    # --------------------------------------------------------
    # UI 평균
    # --------------------------------------------------------

    def get_ui_averages(self):

        with self._buffer_lock:

            pm10_buffer = list(
                self.data_buffer["pm10"]
            )

            pm25_buffer = list(
                self.data_buffer["pm25"]
            )

            pm1_buffer = list(
                self.data_buffer["pm1"]
            )

        if (
            not pm10_buffer
            or not pm25_buffer
            or not pm1_buffer
        ):
            return None

        return (
            int(
                sum(pm10_buffer)
                / len(pm10_buffer)
            ),
            int(
                sum(pm25_buffer)
                / len(pm25_buffer)
            ),
            int(
                sum(pm1_buffer)
                / len(pm1_buffer)
            ),
        )

    # --------------------------------------------------------
    # Serial
    # --------------------------------------------------------

    def run(self):

        is_connected = False

        no_data_error_sent = False
        out_of_range_error_sent = False

        last_ui_update_time = 0.0
        last_data_time = time.time()

        while not self.isInterruptionRequested():

            ser = None

            try:

                ser = serial.Serial(
                    self.port_name,
                    Config.BAUD_RATE,
                    timeout=Config.SERIAL_TIMEOUT
                )

                ser.reset_input_buffer()

                if not is_connected:

                    self.logger.write_error(
                        f"통신 연결 성공 "
                        f"({self.port_name})"
                    )

                is_connected = True

                no_data_error_sent = False
                out_of_range_error_sent = False

                last_data_time = time.time()

                self.current_minute_bucket = (
                    self._get_current_minute()
                )

                self.clear_ui_buffers()

                while (
                    ser.is_open
                    and not self.isInterruptionRequested()
                ):

                    parsed_values = []

                    # ----------------------------------------
                    # 현재 들어와 있는 데이터 처리
                    # ----------------------------------------

                    while ser.in_waiting > 0:

                        raw_bytes = ser.readline()

                        raw_data = raw_bytes.decode(
                            "utf-8",
                            errors="ignore"
                        ).strip()

                        if not raw_data:
                            continue

                        parsed = DustParser.parse(
                            raw_data,
                            self.calib_params
                        )

                        if parsed == "OUT_OF_RANGE":

                            if not out_of_range_error_sent:

                                self.error_signal.emit(
                                    self.sensor_index,
                                    "측정값 범위 초과"
                                )

                                self.logger.write_error(
                                    f"범위 초과 Raw: "
                                    f"{raw_data}"
                                )

                                out_of_range_error_sent = True

                            continue

                        if parsed is not None:

                            parsed_values.append(
                                parsed
                            )

                            last_data_time = (
                                time.time()
                            )

                            out_of_range_error_sent = False

                    # ----------------------------------------
                    # 분 경계
                    # ----------------------------------------

                    self.check_minute_boundary()

                    # ----------------------------------------
                    # 데이터 없음
                    # ----------------------------------------

                    if (
                        time.time()
                        - last_data_time
                        >= Config.NO_DATA_TIMEOUT
                    ):

                        if not no_data_error_sent:

                            self.error_signal.emit(
                                self.sensor_index,
                                "데이터 없음"
                            )

                            self.logger.write_error(
                                f"{Config.NO_DATA_TIMEOUT:.0f}"
                                f"초 이상 데이터 없음"
                            )

                            no_data_error_sent = True

                            self.clear_ui_buffers()

                    elif parsed_values:

                        no_data_error_sent = False

                        # ------------------------------------
                        # 보정된 값 누적
                        # ------------------------------------

                        for parsed in parsed_values:

                            self.add_measurement(
                                parsed["pm10"],
                                parsed["pm25"],
                                parsed["pm1"]
                            )

                        # ------------------------------------
                        # UI 1초 업데이트
                        # ------------------------------------

                        current_time = time.time()

                        if (
                            current_time
                            - last_ui_update_time
                            >= 1.0
                        ):

                            last_ui_update_time = (
                                current_time
                            )

                            averages = (
                                self.get_ui_averages()
                            )

                            if averages is not None:

                                self.data_signal.emit(
                                    self.sensor_index,
                                    averages[0],
                                    averages[1],
                                    averages[2]
                                )

                    self.safe_sleep(100)

            except serial.SerialException as e:

                if is_connected:

                    self.logger.write_error(
                        f"시리얼 통신 오류: {e}"
                    )

                is_connected = False

                self.clear_all_buffers()

                self.error_signal.emit(
                    self.sensor_index,
                    "연결 실패/끊김"
                )

            except Exception:

                if is_connected:

                    self.logger.write_error(
                        "시스템 오류:\n"
                        + traceback.format_exc()
                    )

                is_connected = False

                self.clear_all_buffers()

                self.error_signal.emit(
                    self.sensor_index,
                    "시스템 오류"
                )

            finally:

                if ser is not None:

                    try:

                        if ser.is_open:
                            ser.close()

                    except Exception:
                        pass

                if not self.isInterruptionRequested():

                    self.safe_sleep(
                        Config.RECONNECT_DELAY_MS
                    )

        # ====================================================
        # 프로그램 종료
        #
        # 마지막 분 데이터는 반드시 동기 저장
        # ====================================================

        try:

            self.save_current_minute_average(
                self.current_minute_bucket,
                synchronous=True
            )

        except Exception as e:

            self.logger.write_error(
                f"Thread 종료 시 "
                f"마지막 데이터 저장 오류: {e}"
            )

    # --------------------------------------------------------
    # 보정값 변경
    # --------------------------------------------------------

    def update_calibration(
        self,
        new_calib_params
    ):

        try:

            new_scale = tuple(
                float(x)
                for x in new_calib_params["scale"]
            )

            new_offset = tuple(
                float(x)
                for x in new_calib_params["offset"]
            )

            if len(new_scale) != 3:
                return False

            if len(new_offset) != 3:
                return False

            if not all(
                math.isfinite(x)
                for x in new_scale
            ):
                return False

            if not all(
                math.isfinite(x)
                for x in new_offset
            ):
                return False

            if any(
                x < 0
                for x in new_scale
            ):
                return False

            self.calib_params = {
                "scale": new_scale,
                "offset": new_offset
            }

            # 보정 변경 후 이전 평균과 섞이지 않게
            # 현재 UI/분 버퍼를 초기화
            self.clear_all_buffers()

            return True

        except Exception:

            return False


# ============================================================
# Dust Level Widget
# ============================================================

class DustLevelWidget(QWidget):

    def __init__(
        self,
        title="미세먼지",
        levels=None,
        parent=None
    ):

        super().__init__(parent)

        self.title = title

        self.levels = (
            levels
            if levels
            else Config.PM10_LEVELS
        )

        self.current_value = 0

        self.initUI()

    def initUI(self):

        self.setStyleSheet(
            "background-color: transparent;"
            "border: none;"
        )

        main_layout = QVBoxLayout(self)

        main_layout.setContentsMargins(
            10,
            10,
            10,
            10
        )

        main_layout.setSpacing(3)

        self.title_label = QLabel(
            self.title
        )

        self.title_label.setFont(
            QFont(
                "Malgun Gothic",
                11,
                QFont.Bold
            )
        )

        self.title_label.setAlignment(
            Qt.AlignCenter
        )

        self.title_label.setStyleSheet(
            "color: black;"
            "border: none;"
        )

        main_layout.addWidget(
            self.title_label
        )

        value_layout = QHBoxLayout()

        value_layout.setAlignment(
            Qt.AlignCenter
        )

        self.val_label = QLabel(
            "----"
        )

        self.val_label.setFont(
            QFont(
                "Arial",
                28,
                QFont.Bold
            )
        )

        self.val_label.setStyleSheet(
            "color: black;"
            "border: none;"
        )

        value_layout.addWidget(
            self.val_label
        )

        unit_label = QLabel(
            "μg/m³"
        )

        unit_label.setFont(
            QFont(
                "Arial",
                10,
                QFont.Bold
            )
        )

        unit_label.setStyleSheet(
            "color: #666666;"
            "margin-bottom: 4px;"
            "border: none;"
        )

        unit_label.setAlignment(
            Qt.AlignBottom
        )

        value_layout.addWidget(
            unit_label
        )

        main_layout.addLayout(
            value_layout
        )

        self.arrow_container = QWidget()

        self.arrow_container.setStyleSheet(
            "border: none;"
        )

        self.arrow_container.setFixedHeight(
            14
        )

        self.arrow_label = QLabel(
            "▼"
        )

        self.arrow_label.setFont(
            QFont(
                "Arial",
                9,
                QFont.Bold
            )
        )

        self.arrow_label.setStyleSheet(
            "color: black;"
            "border: none;"
        )

        self.arrow_label.setAlignment(
            Qt.AlignCenter
        )

        self.arrow_label.setFixedWidth(
            14
        )

        self.arrow_label.setParent(
            self.arrow_container
        )

        main_layout.addWidget(
            self.arrow_container
        )

        level_bar_frame = QFrame()

        level_bar_frame.setStyleSheet(
            "border: none;"
        )

        self.level_bar_layout = QHBoxLayout(
            level_bar_frame
        )

        self.level_bar_layout.setContentsMargins(
            0,
            0,
            0,
            0
        )

        self.level_bar_layout.setSpacing(
            2
        )

        self.level_bars = []

        for i, level in enumerate(
            self.levels
        ):

            bar = QFrame()

            bar.setFixedHeight(
                10
            )

            border_radius = ""

            if i == 0:

                border_radius = (
                    "border-top-left-radius: 5px;"
                    "border-bottom-left-radius: 5px;"
                )

            elif i == len(self.levels) - 1:

                border_radius = (
                    "border-top-right-radius: 5px;"
                    "border-bottom-right-radius: 5px;"
                )

            bar.setStyleSheet(
                f"background-color: "
                f"{level['color']}; "
                f"{border_radius}"
            )

            self.level_bars.append(
                bar
            )

            self.level_bar_layout.addWidget(
                bar
            )

        main_layout.addWidget(
            level_bar_frame
        )

        self.status_text_label = QLabel(
            "대기 중..."
        )

        self.status_text_label.setFont(
            QFont(
                "Malgun Gothic",
                11,
                QFont.Bold
            )
        )

        self.status_text_label.setAlignment(
            Qt.AlignCenter
        )

        self.status_text_label.setFixedHeight(
            30
        )

        self.status_text_label.setStyleSheet(
            "background-color: #E0E0E0;"
            "border-radius: 6px;"
            "color: gray;"
        )

        main_layout.addWidget(
            self.status_text_label
        )

    def update_val(
        self,
        value
    ):

        self.arrow_label.show()

        self.current_value = value

        self.val_label.setText(
            str(value)
        )

        current_level = self.levels[
            self._find_level_index(value)
        ]

        self.status_text_label.setText(
            current_level["name"]
        )

        text_color = (
            "black"
            if current_level["name"] == "보통"
            else "white"
        )

        self.status_text_label.setStyleSheet(
            f"background-color: "
            f"{current_level['color']};"
            f"border-radius: 6px;"
            f"color: {text_color};"
            f"border: none;"
        )

        self.update_arrow_position()

    def _find_level_index(
        self,
        value
    ):

        for i, level in enumerate(
            self.levels
        ):

            if value <= level["max"]:
                return i

        return len(self.levels) - 1

    def set_error_state(
        self,
        msg
    ):

        self.val_label.setText(
            "----"
        )

        self.status_text_label.setText(
            msg
        )

        self.status_text_label.setStyleSheet(
            "background-color: #FFCDD2;"
            "border-radius: 6px;"
            "color: #B71C1C;"
            "border: none;"
        )

        self.current_value = (
            self.levels[0]["min"]
        )

        self.update_arrow_position()

        self.arrow_label.hide()

    def update_arrow_position(self):

        if (
            not self.level_bars
            or self.level_bars[0].geometry().width() == 0
        ):
            return

        value = self.current_value

        level_index = (
            self._find_level_index(value)
        )

        current_level = (
            self.levels[level_index]
        )

        level_min = current_level["min"]
        level_max = current_level["max"]

        if value <= level_min:

            ratio = 0.0

        elif value >= level_max:

            ratio = 1.0

        else:

            denominator = (
                level_max - level_min
            )

            ratio = (
                0.0
                if denominator <= 0
                else (
                    value - level_min
                ) / denominator
            )

        target_bar = (
            self.level_bars[level_index]
        )

        target_x = (
            target_bar.geometry().x()
            + (
                target_bar.geometry().width()
                * ratio
            )
        )

        arrow_x = int(
            target_x
            - (
                self.arrow_label.width()
                / 2
            )
        )

        max_x = (
            self.arrow_container.width()
            - self.arrow_label.width()
        )

        arrow_x = max(
            0,
            min(
                arrow_x,
                max_x
            )
        )

        self.arrow_label.move(
            QPoint(
                arrow_x,
                int(
                    self.arrow_container.height()
                    - self.arrow_label.height()
                    + 2
                )
            )
        )

    def resizeEvent(
        self,
        event
    ):

        super().resizeEvent(
            event
        )

        self.update_arrow_position()


# ============================================================
# Port Selection Dialog
# ============================================================

class PortSelectionDialog(QDialog):

    def __init__(
        self,
        parent=None
    ):

        super().__init__(parent)

        self.selected_ports = {}

        self.available_ports = sorted(
            [
                port.device
                for port
                in serial.tools.list_ports.comports()
            ]
        )

        self.initUI()

    def initUI(self):

        self.setWindowTitle(
            "센서 포트 설정"
        )

        self.resize(
            350,
            250
        )

        self.setStyleSheet(
            "background-color: white;"
        )

        layout = QVBoxLayout(
            self
        )

        layout.setContentsMargins(
            20,
            20,
            20,
            20
        )

        layout.setSpacing(
            15
        )

        title_label = QLabel(
            "모니터링할 센서 포트를 선택하세요"
        )

        title_label.setFont(
            QFont(
                "Malgun Gothic",
                10,
                QFont.Bold
            )
        )

        layout.addWidget(
            title_label
        )

        self.combos = []

        for i in range(
            Config.MAX_SENSORS
        ):

            h_layout = QHBoxLayout()

            label = QLabel(
                f"센서 {i + 1} 포트:"
            )

            label.setFont(
                QFont(
                    "Malgun Gothic",
                    9
                )
            )

            combo = QComboBox()

            combo.addItem(
                "선택 안 함"
            )

            combo.addItems(
                self.available_ports
            )

            combo.setFont(
                QFont(
                    "Malgun Gothic",
                    9
                )
            )

            combo.currentIndexChanged.connect(
                self.update_combo_items
            )

            h_layout.addWidget(
                label
            )

            h_layout.addWidget(
                combo
            )

            layout.addLayout(
                h_layout
            )

            self.combos.append(
                combo
            )

        self.start_btn = QPushButton(
            "모니터링 시작"
        )

        self.start_btn.setFixedHeight(
            35
        )

        self.start_btn.setFont(
            QFont(
                "Malgun Gothic",
                10,
                QFont.Bold
            )
        )

        self.start_btn.setStyleSheet(
            "background-color: #007BFF;"
            "color: white;"
            "border-radius: 4px;"
        )

        self.start_btn.clicked.connect(
            self.accept_selection
        )

        layout.addWidget(
            self.start_btn
        )

    def update_combo_items(self):

        selected_values = [
            combo.currentText()
            for combo in self.combos
            if (
                combo.currentText()
                and combo.currentText()
                != "선택 안 함"
            )
        ]

        for combo in self.combos:

            current_selection = (
                combo.currentText()
            )

            combo.blockSignals(
                True
            )

            combo.clear()

            combo.addItem(
                "선택 안 함"
            )

            for port in self.available_ports:

                if (
                    port in selected_values
                    and port != current_selection
                ):
                    continue

                combo.addItem(
                    port
                )

            index = combo.findText(
                current_selection
            )

            combo.setCurrentIndex(
                index
                if index != -1
                else 0
            )

            combo.blockSignals(
                False
            )

    def accept_selection(self):

        self.selected_ports = {
            i: combo.currentText()
            for i, combo in enumerate(
                self.combos
            )
            if (
                combo.currentText()
                and combo.currentText()
                != "선택 안 함"
            )
        }

        if not self.selected_ports:

            QMessageBox.warning(
                self,
                "경고",
                "최소 하나 이상의 센서 포트를 선택해 주세요."
            )

            return

        self.accept()


# ============================================================
# Calibration Dialog
# ============================================================

class CalibrationDialog(QDialog):

    def __init__(
        self,
        current_calib_dict,
        parent=None
    ):

        super().__init__(
            parent
        )

        self.calib_dict = {
            i: {
                "scale": tuple(
                    v["scale"]
                ),
                "offset": tuple(
                    v["offset"]
                )
            }
            for i, v
            in current_calib_dict.items()
        }

        self.initUI()

    def initUI(self):

        self.setWindowTitle(
            "센서 보정값 설정"
        )

        self.resize(
            500,
            500
        )

        self.setStyleSheet(
            "background-color: white;"
        )

        main_layout = QVBoxLayout(
            self
        )

        main_layout.setContentsMargins(
            20,
            20,
            20,
            20
        )

        main_layout.setSpacing(
            15
        )

        title_label = QLabel(
            "센서별 PM10, PM2.5, PM1.0 보정 계수 설정"
        )

        title_label.setFont(
            QFont(
                "Malgun Gothic",
                10,
                QFont.Bold
            )
        )

        main_layout.addWidget(
            title_label
        )

        formula_label = QLabel(
            "📌 계산 공식: 결과 값 = (원본값 × Scale) + Offset"
        )

        formula_label.setFont(
            QFont(
                "Malgun Gothic",
                9
            )
        )

        formula_label.setStyleSheet(
            "color: #0056b3;"
            "background-color: #e7f1ff;"
            "padding: 6px;"
            "border-radius: 4px;"
        )

        main_layout.addWidget(
            formula_label
        )

        scroll = QScrollArea()

        scroll.setWidgetResizable(
            True
        )

        scroll_content = QWidget()

        scroll_layout = QVBoxLayout(
            scroll_content
        )

        scroll_layout.setSpacing(
            10
        )

        self.inputs = {}

        for i in range(
            Config.MAX_SENSORS
        ):

            group = QGroupBox(
                f"센서 {i + 1}"
            )

            group.setFont(
                QFont(
                    "Malgun Gothic",
                    9,
                    QFont.Bold
                )
            )

            g_layout = QVBoxLayout(
                group
            )

            g_layout.setSpacing(
                6
            )

            curr = self.calib_dict.get(
                i,
                {
                    "scale": (
                        1.0,
                        1.0,
                        1.0
                    ),
                    "offset": (
                        0.0,
                        0.0,
                        0.0
                    )
                }
            )

            scales = curr["scale"]
            offsets = curr["offset"]

            self.inputs[i] = {
                "s_10": QLineEdit(
                    str(scales[0])
                ),
                "o_10": QLineEdit(
                    str(offsets[0])
                ),

                "s_25": QLineEdit(
                    str(scales[1])
                ),
                "o_25": QLineEdit(
                    str(offsets[1])
                ),

                "s_1": QLineEdit(
                    str(scales[2])
                ),
                "o_1": QLineEdit(
                    str(offsets[2])
                ),
            }

            rows = [
                (
                    "PM10",
                    "s_10",
                    "o_10"
                ),
                (
                    "PM2.5",
                    "s_25",
                    "o_25"
                ),
                (
                    "PM1.0",
                    "s_1",
                    "o_1"
                ),
            ]

            for (
                target_name,
                key_s,
                key_o
            ) in rows:

                row_layout = QHBoxLayout()

                lbl = QLabel(
                    f"• {target_name}"
                )

                lbl.setFixedWidth(
                    50
                )

                lbl.setFont(
                    QFont(
                        "Malgun Gothic",
                        9
                    )
                )

                row_layout.addWidget(
                    lbl
                )

                row_layout.addWidget(
                    QLabel("Scale:")
                )

                row_layout.addWidget(
                    self.inputs[i][key_s]
                )

                row_layout.addWidget(
                    QLabel("Offset:")
                )

                row_layout.addWidget(
                    self.inputs[i][key_o]
                )

                g_layout.addLayout(
                    row_layout
                )

            scroll_layout.addWidget(
                group
            )

        scroll.setWidget(
            scroll_content
        )

        main_layout.addWidget(
            scroll
        )

        save_btn = QPushButton(
            "적용 및 저장"
        )

        save_btn.setFixedHeight(
            35
        )

        save_btn.setFont(
            QFont(
                "Malgun Gothic",
                10,
                QFont.Bold
            )
        )

        save_btn.setStyleSheet(
            "background-color: #28A745;"
            "color: white;"
            "border-radius: 4px;"
        )

        save_btn.clicked.connect(
            self.save_calib
        )

        main_layout.addWidget(
            save_btn
        )

    def save_calib(self):

        try:

            new_calib = {}

            for i in range(
                Config.MAX_SENSORS
            ):

                s10 = float(
                    self.inputs[i]["s_10"].text()
                )

                o10 = float(
                    self.inputs[i]["o_10"].text()
                )

                s25 = float(
                    self.inputs[i]["s_25"].text()
                )

                o25 = float(
                    self.inputs[i]["o_25"].text()
                )

                s1 = float(
                    self.inputs[i]["s_1"].text()
                )

                o1 = float(
                    self.inputs[i]["o_1"].text()
                )

                values = (
                    s10,
                    o10,
                    s25,
                    o25,
                    s1,
                    o1
                )

                # NaN / Inf 차단
                if not all(
                    math.isfinite(x)
                    for x in values
                ):
                    raise ValueError

                # Scale 음수 방지
                if (
                    s10 < 0
                    or s25 < 0
                    or s1 < 0
                ):
                    QMessageBox.warning(
                        self,
                        "오류",
                        "Scale 값은 0 이상이어야 합니다."
                    )
                    return

                new_calib[i] = {
                    "scale": (
                        s10,
                        s25,
                        s1
                    ),
                    "offset": (
                        o10,
                        o25,
                        o1
                    )
                }

            self.calib_dict = new_calib

            self.accept()

        except ValueError:

            QMessageBox.warning(
                self,
                "오류",
                "모든 입력값에는 올바른 숫자(실수)를 입력해주세요."
            )


# ============================================================
# Excel Export Worker
# ============================================================

class ExcelExportWorker(QThread):

    finished_signal = pyqtSignal(
        bool,
        str
    )

    def __init__(
        self,
        logger,
        date_str,
        parent=None
    ):

        super().__init__(
            parent
        )

        self.logger = logger
        self.date_str = date_str

    def run(self):

        try:

            if self.isInterruptionRequested():
                return

            success = self.logger.export_excel(
                self.date_str
            )

            if success:

                self.finished_signal.emit(
                    True,
                    f"{self.date_str} "
                    f"Excel 저장 완료"
                )

            else:

                self.finished_signal.emit(
                    False,
                    f"{self.date_str} "
                    f"저장할 데이터가 없습니다."
                )

        except Exception as e:

            self.finished_signal.emit(
                False,
                f"Excel 저장 오류: {e}"
            )


# ============================================================
# Main Application
# ============================================================

class DustMonitorApp(QMainWindow):

    def __init__(
        self,
        slot_mapping
    ):

        super().__init__()

        self.threads = []

        self.sensor_widgets = {}

        self.slot_mapping = slot_mapping

        self.db_path = os.path.join(
            BASE_DIR,
            "Data",
            "dust_measurement.db"
        )

        # 자동 Excel 작업 관리
        self.export_workers = []

        # 이미 자동 저장한 날짜 추적
        self.exported_dates = set()

        self.initUI()

        self.start_monitoring()

    # --------------------------------------------------------
    # UI
    # --------------------------------------------------------

    def initUI(self):

        self.setWindowTitle(
            "다중 미세먼지 모니터링 시스템"
        )

        self.setStyleSheet(
            "QMainWindow {"
            "background-color: white;"
            "}"
        )

        self.setMinimumSize(
            400,
            150
        )

        central_widget = QWidget()

        main_layout = QVBoxLayout(
            central_widget
        )

        # ----------------------------------------
        # 상단
        # ----------------------------------------

        top_control_layout = QHBoxLayout()

        self.calib_btn = QPushButton(
            "센서 보정 설정"
        )

        self.calib_btn.setFont(
            QFont(
                "Malgun Gothic",
                9,
                QFont.Bold
            )
        )

        self.calib_btn.setFixedHeight(
            30
        )

        self.calib_btn.clicked.connect(
            self.open_calibration_dialog
        )

        top_control_layout.addWidget(
            self.calib_btn
        )

        # ----------------------------------------
        # 기존 "전체 엑셀 내보내기" 버튼 제거
        # ----------------------------------------

        top_control_layout.addStretch(
            1
        )

        main_layout.addLayout(
            top_control_layout
        )

        # ----------------------------------------
        # 센서 영역
        # ----------------------------------------

        scroll = QScrollArea()

        scroll.setWidgetResizable(
            True
        )

        scroll_content = QWidget()

        scroll.setWidget(
            scroll_content
        )

        self.scroll_layout = QVBoxLayout(
            scroll_content
        )

        self.scroll_layout.setContentsMargins(
            15,
            15,
            15,
            15
        )

        self.scroll_layout.setSpacing(
            15
        )

        self.scroll_layout.addStretch(
            1
        )

        main_layout.addWidget(
            scroll
        )

        self.setCentralWidget(
            central_widget
        )

    # --------------------------------------------------------
    # Calibration
    # --------------------------------------------------------

    def open_calibration_dialog(self):

        dialog = CalibrationDialog(
            Config.SENSOR_CALIBRATION,
            self
        )

        if (
            dialog.exec_()
            == QDialog.Accepted
        ):

            # ----------------------------------------------
            # 1. Config 변경
            # ----------------------------------------------

            Config.SENSOR_CALIBRATION = (
                dialog.calib_dict
            )

            # ----------------------------------------------
            # 2. 파일 저장
            # ----------------------------------------------

            saved = Config.save_calibration()

            if not saved:

                QMessageBox.warning(
                    self,
                    "오류",
                    "보정값 파일 저장에 실패했습니다."
                )

                return

            # ----------------------------------------------
            # 3. 현재 실행 중인 Thread에 즉시 적용
            # ----------------------------------------------

            apply_failed = []

            for thread in self.threads:

                idx = thread.sensor_index

                calib = (
                    Config.SENSOR_CALIBRATION.get(
                        idx
                    )
                )

                if calib is not None:

                    if not thread.update_calibration(
                        calib
                    ):

                        apply_failed.append(
                            idx + 1
                        )

            if apply_failed:

                QMessageBox.warning(
                    self,
                    "경고",
                    "보정값 파일은 저장되었지만 "
                    "일부 센서에 적용하지 못했습니다.\n"
                    f"센서: {apply_failed}"
                )

            else:

                QMessageBox.information(
                    self,
                    "성공",
                    "보정값이 저장되고 "
                    "현재 센서에 즉시 반영되었습니다."
                )

    # --------------------------------------------------------
    # Monitoring 시작
    # --------------------------------------------------------

    def start_monitoring(self):

        for (
            sensor_index,
            port_name
        ) in self.slot_mapping.items():

            group_box = QGroupBox(
                f"센서 {sensor_index + 1} "
                f"포트: {port_name}"
            )

            group_box.setFont(
                QFont(
                    "Malgun Gothic",
                    11,
                    QFont.Bold
                )
            )

            port_layout = QHBoxLayout(
                group_box
            )

            port_layout.setSpacing(
                10
            )

            w_pm10 = DustLevelWidget(
                "PM10 (미세먼지)",
                Config.PM10_LEVELS
            )

            w_pm25 = DustLevelWidget(
                "PM2.5 (초미세먼지)",
                Config.PM25_LEVELS
            )

            w_pm1 = DustLevelWidget(
                "PM1.0 (극초미세먼지)",
                Config.PM1_LEVELS
            )

            port_layout.addWidget(
                w_pm10
            )

            port_layout.addWidget(
                w_pm25
            )

            port_layout.addWidget(
                w_pm1
            )

            self.scroll_layout.insertWidget(
                self.scroll_layout.count() - 1,
                group_box
            )

            self.sensor_widgets[
                sensor_index
            ] = {
                "pm10": w_pm10,
                "pm25": w_pm25,
                "pm1": w_pm1
            }

            thread = SerialThread(
                port_name,
                sensor_index
            )

            thread.data_signal.connect(
                self.update_data
            )

            thread.error_signal.connect(
                self.handle_error
            )

            # ----------------------------------------------
            # 날짜 변경 이벤트
            # ----------------------------------------------

            thread.date_changed_signal.connect(
                self.handle_date_changed
            )

            thread.start()

            self.threads.append(
                thread
            )

        sensor_count = len(
            self.slot_mapping
        )

        self.resize(
            980,
            max(
                220,
                sensor_count * 245 + 60
            )
        )

    # --------------------------------------------------------
    # 실시간 데이터
    # --------------------------------------------------------

    def update_data(
        self,
        sensor_index,
        pm10,
        pm25,
        pm1
    ):

        widgets = self.sensor_widgets.get(
            sensor_index
        )

        if not widgets:
            return

        widgets["pm10"].update_val(
            pm10
        )

        widgets["pm25"].update_val(
            pm25
        )

        widgets["pm1"].update_val(
            pm1
        )

    # --------------------------------------------------------
    # 오류
    # --------------------------------------------------------

    def handle_error(
        self,
        sensor_index,
        error_msg
    ):

        widgets = self.sensor_widgets.get(
            sensor_index
        )

        if not widgets:
            return

        widgets["pm10"].set_error_state(
            error_msg
        )

        widgets["pm25"].set_error_state(
            error_msg
        )

        widgets["pm1"].set_error_state(
            error_msg
        )

    # ========================================================
    # 날짜 변경 -> 이전 날짜 Excel 자동 저장
    # ========================================================

    def handle_date_changed(
        self,
        sensor_index,
        previous_date
    ):

        # 같은 센서/날짜 중복 요청 방지
        key = (
            sensor_index,
            previous_date
        )

        if key in self.exported_dates:
            return

        self.exported_dates.add(
            key
        )

        # 해당 Thread 찾기
        target_thread = None

        for thread in self.threads:

            if (
                thread.sensor_index
                == sensor_index
            ):

                target_thread = thread
                break

        if target_thread is None:
            return

        # ----------------------------------------------
        # DB queue에서 이전 분 저장이 먼저 처리되도록
        # flush 후 Excel export
        # ----------------------------------------------

        db_manager = (
            target_thread.logger.db_manager
        )

        db_manager.flush(
            timeout=10.0
        )

        worker = ExcelExportWorker(
            target_thread.logger,
            previous_date,
            self
        )

        worker.finished_signal.connect(
            self.handle_auto_export_finished
        )

        # worker가 GC되지 않도록 보관
        self.export_workers.append(
            worker
        )

        worker.finished.connect(
            lambda w=worker:
            self.remove_export_worker(w)
        )

        worker.start()

    def remove_export_worker(
        self,
        worker
    ):

        try:

            if worker in self.export_workers:
                self.export_workers.remove(
                    worker
                )

        except Exception:
            pass

    def handle_auto_export_finished(
        self,
        success,
        message
    ):

        # 자동 저장은 팝업을 띄우지 않는다.
        #
        # 실패한 경우에만 콘솔에 기록.
        if not success:
            print(
                f"[자동 Excel] {message}"
            )

    # ========================================================
    # 프로그램 종료
    # ========================================================

    def closeEvent(
        self,
        event
    ):

        # ----------------------------------------
        # 1. 자동 Excel Worker 종료 요청
        # ----------------------------------------

        for worker in list(
            self.export_workers
        ):

            if worker.isRunning():

                worker.requestInterruption()

                worker.wait(
                    2000
                )

        # ----------------------------------------
        # 2. Serial Thread 중단 요청
        # ----------------------------------------

        for thread in self.threads:

            if thread.isRunning():

                thread.requestInterruption()

        # ----------------------------------------
        # 3. Serial Thread 종료 대기
        #
        # Thread 종료 시 마지막 분 데이터를
        # 동기 저장한다.
        # ----------------------------------------

        all_stopped = True

        for thread in self.threads:

            if thread.isRunning():

                thread.wait(
                    Config.THREAD_WAIT_MS
                )

            if thread.isRunning():

                all_stopped = False

                print(
                    f"센서 {thread.sensor_index + 1} "
                    f"Thread 강제 종료"
                )

                thread.terminate()

                thread.wait(
                    1000
                )

        # ----------------------------------------
        # 4. DB queue flush
        # ----------------------------------------

        db_manager = DatabaseManager(
            self.db_path
        )

        db_manager.flush(
            timeout=10.0
        )

        # ----------------------------------------
        # 5. 당일 Excel 자동 저장
        #
        # 버튼 없음.
        # 프로그램 종료 시 무조건 수행.
        # ----------------------------------------

        today = datetime.now().strftime(
            "%Y-%m-%d"
        )

        for thread in self.threads:

            try:

                # DB queue에 남은 저장을 먼저 처리
                db_manager.flush(
                    timeout=10.0
                )

                success = (
                    thread.logger.export_excel(
                        today
                    )
                )

                if not success:

                    print(
                        f"센서 {thread.sensor_index + 1} "
                        f"당일 Excel 저장: "
                        f"데이터 없음 또는 실패"
                    )

            except Exception as e:

                thread.logger.write_error(
                    f"종료 Excel 처리 오류: {e}"
                )

        # ----------------------------------------
        # 6. Excel 저장 완료 후 DB Worker 종료
        # ----------------------------------------

        db_manager.stop(
            timeout=10.0
        )

        # ----------------------------------------
        # 7. Thread가 비정상 종료되었다면 경고
        # ----------------------------------------

        if not all_stopped:

            QMessageBox.warning(
                self,
                "종료 경고",
                "일부 센서 스레드가 "
                "정상 종료되지 않았습니다."
            )

        event.accept()

    # --------------------------------------------------------
    # 마우스로 창 이동
    # --------------------------------------------------------

    def mousePressEvent(
        self,
        event
    ):

        if event.button() == Qt.LeftButton:

            self.dragPos = (
                event.globalPos()
            )

    def mouseMoveEvent(
        self,
        event
    ):

        if (
            event.buttons()
            == Qt.LeftButton
            and hasattr(
                self,
                "dragPos"
            )
        ):

            self.move(
                self.pos()
                + event.globalPos()
                - self.dragPos
            )

            self.dragPos = (
                event.globalPos()
            )

    # --------------------------------------------------------
    # ESC 종료
    # --------------------------------------------------------

    def keyPressEvent(
        self,
        event
    ):

        if event.key() == Qt.Key_Escape:

            self.close()

        else:

            super().keyPressEvent(
                event
            )


# ============================================================
# Main
# ============================================================

def main():

    app = QApplication(
        sys.argv
    )

    # ----------------------------------------
    # 보정값 로드
    # ----------------------------------------

    Config.load_calibration()

    # ----------------------------------------
    # DB 초기화
    # ----------------------------------------

    db_path = os.path.join(
        BASE_DIR,
        "Data",
        "dust_measurement.db"
    )

    SensorLogger.init_global_database(
        db_path
    )

    # ----------------------------------------
    # 오래된 로그 정리
    # ----------------------------------------

    SensorLogger.cleanup_old_logs_global()

    # ----------------------------------------
    # 포트 선택
    # ----------------------------------------

    dialog = PortSelectionDialog()

    if (
        dialog.exec_()
        != QDialog.Accepted
    ):

        return 0

    # ----------------------------------------
    # 메인 화면
    # ----------------------------------------

    main_window = DustMonitorApp(
        dialog.selected_ports
    )

    main_window.show()

    return app.exec_()


# ============================================================
# Entry
# ============================================================

if __name__ == "__main__":

    sys.exit(
        main()
    )