import sqlite3
import csv
from datetime import datetime, timedelta
import sys
import os
import serial
import time
import re
import traceback
import serial.tools.list_ports

from PyQt5.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout,
    QHBoxLayout, QLabel, QFrame, QComboBox,
    QPushButton, QDialog, QMessageBox, QMenu, QAction
)
from PyQt5.QtCore import QThread, pyqtSignal, Qt, QPoint, QTimer
from PyQt5.QtGui import QFont

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ============================================================
# 1. 기본 경로
# ============================================================

if getattr(sys, "frozen", False):
    BASE_DIR = os.path.dirname(sys.executable)
else:
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))


# ============================================================
# 2. 사용자 설정
# ============================================================

BAUD_RATE = 9600

# 센서별 2점 보정
# 보정값 = 원시값 × SLOPE + OFFSET
SENSOR_CALIB_PARAMS = {
    0: (1.000000, 0.0),
    1: (1.084081, 0.0),
    2: (1.008490, 0.0),
}

# 허용 CO2 범위
CO2_MIN = 0
CO2_MAX = 30000

# 화면 표시 이동평균 개수
SMOOTHING_WINDOW = 1

# 데이터 없음 판단 시간
NO_DATA_TIMEOUT = 30.0

# 재연결 대기
RECONNECT_DELAY_MS = 2000

# 로그 보관 기간
LOG_RETENTION_DAYS = 30


# ============================================================
# 3. 공통 함수
# ============================================================

def get_current_minute():
    """
    현재 시간을 초/마이크로초 제거한 분 단위 datetime으로 반환
    """
    return datetime.now().replace(second=0, microsecond=0)


def safe_int(value, default=0):
    try:
        return int(round(float(value)))
    except (ValueError, TypeError):
        return default


# ============================================================
# 4. 포트 선택 다이얼로그
# ============================================================

class PortSelectionDialog(QDialog):

    def __init__(self, parent=None):
        super().__init__(parent)

        self.setWindowTitle("시리얼 포트 설정")
        self.setFixedSize(360, 280)

        self.combos = []

        self.initUI()

    def initUI(self):

        layout = QVBoxLayout(self)
        layout.setSpacing(12)
        layout.setContentsMargins(20, 20, 20, 20)

        title = QLabel("센서 포트 지정")
        title.setFont(QFont("Malgun Gothic", 12, QFont.Bold))
        title.setAlignment(Qt.AlignCenter)

        layout.addWidget(title)

        available_ports = sorted(
            [port.device for port in serial.tools.list_ports.comports()]
        )

        if not available_ports:
            available_ports = ["포트 없음"]

        for i in range(3):

            row = QHBoxLayout()

            label = QLabel(f"센서 {i + 1}:")
            label.setFont(QFont("Malgun Gothic", 10))

            combo = QComboBox()
            combo.setFont(QFont("Malgun Gothic", 10))

            combo.addItem("선택 안 함")

            for port in available_ports:
                if port != "포트 없음":
                    combo.addItem(port)

            combo.currentIndexChanged.connect(self.update_port_lists)

            self.combos.append(combo)

            row.addWidget(label)
            row.addWidget(combo)

            layout.addLayout(row)

        start_button = QPushButton("모니터링 시작")

        start_button.setFont(
            QFont("Malgun Gothic", 10, QFont.Bold)
        )

        start_button.setStyleSheet("""
            QPushButton {
                padding: 8px;
                background-color: #007BFF;
                color: white;
                border-radius: 4px;
            }

            QPushButton:hover {
                background-color: #0056b3;
            }
        """)

        start_button.clicked.connect(self.on_start)

        layout.addWidget(start_button)

    def update_port_lists(self):

        selected_ports = []

        for combo in self.combos:

            port = combo.currentText()

            if port not in ("", "선택 안 함", "포트 없음"):
                selected_ports.append(port)

        for combo in self.combos:

            current = combo.currentText()

            combo.blockSignals(True)

            combo.clear()

            combo.addItem("선택 안 함")

            for port in sorted(
                [p.device for p in serial.tools.list_ports.comports()]
            ):

                if port in selected_ports and port != current:
                    continue

                combo.addItem(port)

            index = combo.findText(current)

            if index >= 0:
                combo.setCurrentIndex(index)
            else:
                combo.setCurrentIndex(0)

            combo.blockSignals(False)

    def on_start(self):

        selected = self.get_selected_ports()

        if not selected:

            QMessageBox.warning(
                self,
                "경고",
                "최소 1개 이상의 포트를 선택해야 합니다."
            )

            return

        if len(selected) != len(set(selected)):

            QMessageBox.warning(
                self,
                "포트 중복",
                "동일한 COM 포트를 두 개 이상의 센서에 지정할 수 없습니다."
            )

            return

        self.accept()

    def get_selected_ports(self):

        selected = []

        for combo in self.combos:

            text = combo.currentText()

            if text not in ("선택 안 함", "포트 없음", ""):
                selected.append(text)

        return selected


# ============================================================
# 5. SQLite 데이터베이스 관리
# ============================================================

class DatabaseManager:

    def __init__(self):

        self.db_dir = os.path.join(
            BASE_DIR,
            "DB_Logs"
        )

        os.makedirs(
            self.db_dir,
            exist_ok=True
        )

        self.db_path = os.path.join(
            self.db_dir,
            "co2_monitoring.db"
        )

        self.init_db()

    def get_connection(self):

        conn = sqlite3.connect(
            self.db_path,
            timeout=15.0
        )

        conn.execute("PRAGMA journal_mode=WAL;")
        conn.execute("PRAGMA busy_timeout=15000;")

        return conn

    def init_db(self):

        try:

            with self.get_connection() as conn:

                conn.execute("""
                    CREATE TABLE IF NOT EXISTS co2_logs (
                        id INTEGER PRIMARY KEY AUTOINCREMENT,
                        today_str TEXT NOT NULL,
                        time_str TEXT NOT NULL,
                        sensor_index INTEGER NOT NULL,
                        port TEXT NOT NULL,
                        record_value TEXT NOT NULL,
                        created_at TEXT NOT NULL
                    )
                """)

                conn.execute("""
                    CREATE INDEX IF NOT EXISTS
                    idx_co2_logs_date_sensor
                    ON co2_logs(today_str, sensor_index)
                """)

                conn.commit()

        except sqlite3.Error as e:

            print(f"SQLite 초기화 오류: {e}")

    def insert_log(
        self,
        today_str,
        time_str,
        sensor_index,
        port,
        record_value
    ):

        try:

            created_at = datetime.now().strftime(
                "%Y-%m-%d %H:%M:%S"
            )

            with self.get_connection() as conn:

                conn.execute("""
                    INSERT INTO co2_logs
                    (
                        today_str,
                        time_str,
                        sensor_index,
                        port,
                        record_value,
                        created_at
                    )
                    VALUES (?, ?, ?, ?, ?, ?)
                """, (
                    today_str,
                    time_str,
                    sensor_index,
                    port,
                    str(record_value),
                    created_at
                ))

                conn.commit()

            return True

        except sqlite3.Error as e:

            print(f"SQLite 저장 오류: {e}")
            return False

    def cleanup_old_data(self):

        try:

            threshold_date = (
                datetime.now() -
                timedelta(days=LOG_RETENTION_DAYS)
            ).strftime("%Y-%m-%d")

            with self.get_connection() as conn:

                conn.execute(
                    """
                    DELETE FROM co2_logs
                    WHERE today_str < ?
                    """,
                    (threshold_date,)
                )

                conn.commit()

        except sqlite3.Error as e:

            print(f"SQLite 오래된 데이터 삭제 오류: {e}")

    def export_to_excel(self):

        try:

            with self.get_connection() as conn:

                rows = conn.execute("""
                    SELECT
                        today_str,
                        time_str,
                        sensor_index,
                        port,
                        record_value
                    FROM co2_logs
                    ORDER BY
                        today_str,
                        time_str,
                        sensor_index
                """).fetchall()

            if not rows:
                return

            grouped_data = {}

            for row in rows:

                today_str = row[0]
                time_str = row[1]
                sensor_index = row[2]
                port = row[3]
                value = row[4]

                try:

                    if "." in str(value):
                        value = float(value)
                    else:
                        value = int(value)

                except (ValueError, TypeError):
                    pass

                key = (
                    today_str,
                    sensor_index
                )

                if key not in grouped_data:
                    grouped_data[key] = []

                grouped_data[key].append([
                    today_str,
                    time_str,
                    sensor_index + 1,
                    port,
                    value
                ])

            save_dir = os.path.join(
                BASE_DIR,
                "Excel_Logs"
            )

            os.makedirs(
                save_dir,
                exist_ok=True
            )

            self._create_excel_files(
                save_dir,
                grouped_data
            )

        except Exception as e:

            print(f"Excel 내보내기 오류: {e}")

    def _create_excel_files(
        self,
        save_dir,
        grouped_data
    ):

        thin_border = Border(
            left=Side(style="thin", color="D3D3D3"),
            right=Side(style="thin", color="D3D3D3"),
            top=Side(style="thin", color="D3D3D3"),
            bottom=Side(style="thin", color="D3D3D3")
        )

        header_fill = PatternFill(
            start_color="1F4E78",
            end_color="1F4E78",
            fill_type="solid"
        )

        header_font = Font(
            name="맑은 고딕",
            size=11,
            bold=True,
            color="FFFFFF"
        )

        data_font = Font(
            name="맑은 고딕",
            size=10
        )

        center_align = Alignment(
            horizontal="center",
            vertical="center"
        )

        right_align = Alignment(
            horizontal="right",
            vertical="center"
        )

        headers = [
            "측정일자",
            "측정시간",
            "센서번호",
            "포트",
            "CO2 1분 평균(ppm)"
        ]

        for (
            today_str,
            sensor_index
        ), data_rows in grouped_data.items():

            file_path = os.path.join(
                save_dir,
                f"CO2_log_{today_str}_Sensor{sensor_index + 1}.xlsx"
            )

            wb = Workbook()
            ws = wb.active
            ws.title = "측정 데이터"

            ws.append(headers)
            ws.row_dimensions[1].height = 25

            for col_num, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_num)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = center_align
                cell.border = thin_border

            for row_data in data_rows:
                ws.append(row_data)
                current_row = ws.max_row
                ws.row_dimensions[current_row].height = 20

                for col_num, value in enumerate(row_data, 1):
                    cell = ws.cell(row=current_row, column=col_num)
                    cell.font = data_font
                    cell.border = thin_border

                    if col_num <= 4:
                        cell.alignment = center_align
                    else:
                        if isinstance(value, (int, float)):
                            cell.alignment = right_align
                            cell.number_format = "#,##0"
                        else:
                            cell.alignment = center_align

            for column in ws.columns:
                max_length = 0
                column_letter = get_column_letter(column[0].column)

                for cell in column:
                    text = str(cell.value or "")
                    if cell.row == 1:
                        length = len(text.encode("utf-8"))
                    else:
                        length = len(text)
                    max_length = max(max_length, length)

                ws.column_dimensions[column_letter].width = max(
                    max_length + 5,
                    12
                )

            wb.save(file_path)


# ============================================================
# 6. 파일 로그 관리
# ============================================================

class FileLogger:

    def __init__(self, sensor_index, port):
        self.sensor_index = sensor_index
        self.port = port

        self.log_dir = os.path.join(BASE_DIR, "Logs")
        self.csv_dir = os.path.join(BASE_DIR, "CSV_Logs")

        os.makedirs(self.log_dir, exist_ok=True)
        os.makedirs(self.csv_dir, exist_ok=True)

    def write_error(self, message):
        file_path = os.path.join(
            self.log_dir,
            f"error_log_Sensor{self.sensor_index + 1}.txt"
        )
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        try:
            with open(file_path, "a", encoding="utf-8-sig") as file:
                file.write(f"{timestamp} | 포트: {self.port} | {message}\n")
        except OSError:
            pass

    def save_csv(self, today_str, time_str, record_value):
        try:
            file_path = os.path.join(
                self.csv_dir,
                f"CO2_log_{today_str}_Sensor{self.sensor_index + 1}.csv"
            )
            exists = os.path.exists(file_path)

            with open(file_path, "a", newline="", encoding="utf-8-sig") as file:
                writer = csv.writer(file)

                if not exists:
                    writer.writerow([
                        "측정일자", "측정시간", "센서번호", "포트", "CO2 1분 평균(ppm)"
                    ])

                writer.writerow([
                    today_str, time_str, self.sensor_index + 1, self.port, record_value
                ])
        except OSError as e:
            self.write_error(f"CSV 저장 오류: {e}")

    @staticmethod
    def cleanup_old_files():
        threshold = time.time() - LOG_RETENTION_DAYS * 86400
        directories = ["Logs", "CSV_Logs", "Excel_Logs"]

        for directory in directories:
            path = os.path.join(BASE_DIR, directory)
            if not os.path.exists(path):
                continue

            try:
                for filename in os.listdir(path):
                    file_path = os.path.join(path, filename)
                    if not os.path.isfile(file_path):
                        continue
                    try:
                        if os.path.getmtime(file_path) < threshold:
                            os.remove(file_path)
                    except OSError:
                        pass
            except OSError:
                pass


# ============================================================
# 7. 시리얼 통신 스레드
# ============================================================

class SerialThread(QThread):

    data_signal = pyqtSignal(int, int)
    error_signal = pyqtSignal(int, str)

    def __init__(self, port_name, sensor_index, db_manager):
        super().__init__()

        self.current_port = port_name
        self.sensor_index = sensor_index
        self.db_manager = db_manager

        self.logger = FileLogger(sensor_index, port_name)

        # 화면 이동평균 버퍼
        self.data_buffer = []
        self.smoothing_window = max(1, int(SMOOTHING_WINDOW))

        # [최적화 1] 1분 평균 계산을 위한 누적 합계 및 카운터 (메모리 O(1) 최적화)
        self.minute_sum = 0.0
        self.minute_count = 0

        self.current_minute = get_current_minute()
        self.last_recorded_status = "정상"
        self.last_data_time = time.time()

    def parse_data(self, raw_data):
        try:
            raw_data = raw_data.strip()
            if not raw_data:
                return None

            raw_co2 = None

            if re.search(r"[ABC]\s*:", raw_data, re.IGNORECASE):
                prefix = chr(65 + self.sensor_index)
                pattern = rf"{prefix}\s*:\s*([-+]?\d*\.?\d+)"
                match = re.search(pattern, raw_data, re.IGNORECASE)
                if match:
                    raw_co2 = float(match.group(1))

            elif "CO2" in raw_data.upper():
                match = re.search(r"CO2\s*[:=]?\s*([-+]?\d*\.?\d+)", raw_data, re.IGNORECASE)
                if match:
                    raw_co2 = float(match.group(1))

            else:
                numbers = re.findall(r"[-+]?\d*\.?\d+", raw_data)
                if len(numbers) == 1:
                    raw_co2 = float(numbers[0])

            if raw_co2 is None:
                return None

            if raw_co2 < CO2_MIN or raw_co2 > CO2_MAX:
                return "OUT_OF_RANGE"

            return raw_co2

        except (ValueError, TypeError, AttributeError):
            return None

    def calibrate(self, raw_value):
        slope, offset = SENSOR_CALIB_PARAMS.get(
            self.sensor_index, (1.0, 0.0)
        )
        calibrated = (raw_value * slope) + offset
        return max(CO2_MIN, min(CO2_MAX, calibrated))

    def clear_buffers(self):
        self.data_buffer.clear()
        self.minute_sum = 0.0
        self.minute_count = 0

    def safe_sleep(self, milliseconds):
        elapsed = 0
        while elapsed < milliseconds and not self.isInterruptionRequested():
            step = min(50, milliseconds - elapsed)
            self.msleep(step)
            elapsed += step

    def save_minute_average(self, minute_datetime):
        today_str = minute_datetime.strftime("%Y-%m-%d")
        time_str = minute_datetime.strftime("%H:%M:00")

        if self.minute_count > 0:
            record_value = safe_int(self.minute_sum / self.minute_count)
        else:
            record_value = f"Error: {self.last_recorded_status}"

        self.logger.save_csv(today_str, time_str, record_value)
        self.db_manager.insert_log(
            today_str, time_str, self.sensor_index, self.current_port, record_value
        )

        # 버퍼 초기화
        self.minute_sum = 0.0
        self.minute_count = 0

    def process_minute_change(self):
        now_minute = get_current_minute()
        if now_minute == self.current_minute:
            return

        previous_minute = self.current_minute
        self.save_minute_average(previous_minute)
        self.current_minute = now_minute

    def run(self):
        FileLogger.cleanup_old_files()

        is_connected = False
        no_data_error_sent = False
        last_ui_update_time = 0.0

        self.current_minute = get_current_minute()
        self.last_data_time = time.time()

        while not self.isInterruptionRequested():
            ser = None
            try:
                ser = serial.Serial(
                    self.current_port,
                    BAUD_RATE,
                    timeout=0.1
                )
                ser.reset_input_buffer()

                is_connected = True
                no_data_error_sent = False
                self.last_data_time = time.time()
                self.current_minute = get_current_minute()
                self.clear_buffers()
                self.last_recorded_status = "정상"
                self.logger.write_error("통신 연결 성공")

                while ser.is_open and not self.isInterruptionRequested():
                    current_time = time.time()
                    self.process_minute_change()

                    while ser.in_waiting > 0:
                        raw_bytes = ser.readline()
                        if not raw_bytes:
                            continue

                        raw_data = raw_bytes.decode(
                            "utf-8", errors="ignore"
                        ).strip()

                        parsed = self.parse_data(raw_data)

                        if parsed == "OUT_OF_RANGE":
                            error_message = "범위 초과 (위험)"
                            self.error_signal.emit(self.sensor_index, error_message)
                            self.logger.write_error(f"{error_message} Raw={raw_data}")
                            self.last_recorded_status = error_message
                            self.last_data_time = current_time
                            continue

                        if parsed is not None:
                            calibrated_value = self.calibrate(parsed)
                            self.last_data_time = current_time
                            self.last_recorded_status = "정상"
                            no_data_error_sent = False

                            # 화면용 버퍼 처리
                            self.data_buffer.append(calibrated_value)
                            if len(self.data_buffer) > self.smoothing_window:
                                self.data_buffer.pop(0)

                            # [최적화 1 적용] 1분 평균 누적 합산
                            self.minute_sum += calibrated_value
                            self.minute_count += 1

                    if current_time - self.last_data_time >= NO_DATA_TIMEOUT:
                        if not no_data_error_sent:
                            error_message = "데이터 없음"
                            self.error_signal.emit(self.sensor_index, error_message)
                            self.logger.write_error("30초 이상 데이터 없음")
                            self.last_recorded_status = "데이터 없음"
                            no_data_error_sent = True
                            self.data_buffer.clear()

                    elif self.data_buffer:
                        if current_time - last_ui_update_time >= 1.0:
                            last_ui_update_time = current_time
                            average = sum(self.data_buffer) / len(self.data_buffer)
                            display_value = safe_int(average)
                            self.data_signal.emit(self.sensor_index, display_value)

                    self.safe_sleep(50)

            except serial.SerialException as e:
                if is_connected:
                    self.logger.write_error(f"시리얼 연결 오류: {e}")
                is_connected = False
                self.clear_buffers()
                self.last_recorded_status = "연결 끊김"
                self.error_signal.emit(self.sensor_index, "연결 실패/끊김")

            except Exception:
                if is_connected:
                    self.logger.write_error("시스템 예외:\n" + traceback.format_exc())
                is_connected = False
                self.clear_buffers()
                self.last_recorded_status = "시스템 오류"
                self.error_signal.emit(self.sensor_index, "시스템 오류")

            finally:
                if ser is not None:
                    try:
                        if ser.is_open:
                            ser.close()
                    except (serial.SerialException, OSError):
                        pass

                if not self.isInterruptionRequested():
                    self.safe_sleep(RECONNECT_DELAY_MS)

        try:
            if self.minute_count > 0:
                self.save_minute_average(self.current_minute)
        except Exception as e:
            self.logger.write_error(f"종료 직전 데이터 저장 오류: {e}")


# ============================================================
# 8. CO2 표시 위젯
# ============================================================

class CO2LevelWidget(QWidget):

    LEVELS = [
        {"name": "좋음", "min": 1, "max": 500, "color": "#28A745"},
        {"name": "보통", "min": 501, "max": 1000, "color": "#FFD700"},
        {"name": "민감군", "min": 1001, "max": 3000, "color": "#FD7E14"},
        {"name": "나쁨", "min": 3001, "max": 5000, "color": "#DC3545"},
        {"name": "매우 나쁨", "min": 5001, "max": 10000, "color": "#800080"},
        {"name": "위험", "min": 10001, "max": 30000, "color": "#795548"}
    ]

    def __init__(self, sensor_index, title="센서", parent=None):
        super().__init__(parent)

        self.sensor_index = sensor_index
        self.title = title
        self.current_value = 0

        # [최적화 2] 창 크기 변경 시 과도한 이벤트 호출을 막기 위한 디바운싱 타이머
        self.resize_timer = QTimer(self)
        self.resize_timer.setSingleShot(True)
        self.resize_timer.timeout.connect(self.update_arrow_position)

        self.initUI()

    def initUI(self):
        self.setStyleSheet("background-color: transparent; border: none;")

        layout = QVBoxLayout(self)
        layout.setContentsMargins(20, 20, 20, 20)
        layout.setSpacing(5)

        title_label = QLabel(self.title)
        title_label.setFont(QFont("Malgun Gothic", 16, QFont.Bold))
        title_label.setAlignment(Qt.AlignCenter)
        title_label.setStyleSheet("color: black; border: none;")
        layout.addWidget(title_label)

        value_layout = QHBoxLayout()
        value_layout.setAlignment(Qt.AlignCenter)
        value_layout.setSpacing(2)

        self.co2_value_label = QLabel("----")
        self.co2_value_label.setFont(QFont("Arial", 42, QFont.Bold))
        self.co2_value_label.setStyleSheet("color: black; border: none;")
        value_layout.addWidget(self.co2_value_label)

        ppm_label = QLabel("ppm")
        ppm_label.setFont(QFont("Arial", 16, QFont.Bold))
        ppm_label.setStyleSheet("color: black; margin-bottom: 8px; border: none;")
        ppm_label.setAlignment(Qt.AlignBottom)
        value_layout.addWidget(ppm_label)

        layout.addLayout(value_layout)

        self.arrow_container = QWidget()
        self.arrow_container.setStyleSheet("border: none;")
        self.arrow_container.setFixedHeight(20)

        self.arrow_label = QLabel("▼")
        self.arrow_label.setFont(QFont("Arial", 12, QFont.Bold))
        self.arrow_label.setStyleSheet("color: black; border: none;")
        self.arrow_label.setAlignment(Qt.AlignCenter)
        self.arrow_label.setFixedWidth(20)
        self.arrow_label.setParent(self.arrow_container)
        self.arrow_label.hide()

        layout.addWidget(self.arrow_container)

        level_frame = QFrame()
        level_frame.setStyleSheet("border: none;")
        level_layout = QHBoxLayout(level_frame)
        level_layout.setContentsMargins(0, 0, 0, 0)
        level_layout.setSpacing(2)

        self.level_bars = []

        for i, level in enumerate(self.LEVELS):
            bar = QFrame()
            bar.setFixedHeight(14)
            radius = ""

            if i == 0:
                radius = "border-top-left-radius:7px; border-bottom-left-radius:7px;"
            elif i == len(self.LEVELS) - 1:
                radius = "border-top-right-radius:7px; border-bottom-right-radius:7px;"

            bar.setStyleSheet(f"background-color: {level['color']}; {radius}")
            self.level_bars.append(bar)
            level_layout.addWidget(bar)

        layout.addWidget(level_frame)

        self.status_text_label = QLabel("대기 중...")
        self.status_text_label.setFont(QFont("Malgun Gothic", 16, QFont.Bold))
        self.status_text_label.setAlignment(Qt.AlignCenter)
        self.status_text_label.setFixedHeight(45)
        self.status_text_label.setStyleSheet("background-color:#E0E0E0; border-radius:8px; color:gray;")
        layout.addWidget(self.status_text_label)

        layout.addStretch(1)

    def get_level_index(self, value):
        for i, level in enumerate(self.LEVELS):
            if value <= level["max"]:
                return i
        return len(self.LEVELS) - 1

    def update_co2(self, value):
        self.current_value = value
        self.co2_value_label.setText(str(value))
        self.arrow_label.show()

        index = self.get_level_index(value)
        level = self.LEVELS[index]

        self.status_text_label.setText(level["name"])
        text_color = "black" if level["name"] == "보통" else "white"
        self.status_text_label.setStyleSheet(
            f"background-color:{level['color']}; border-radius:8px; color:{text_color};"
        )

        self.update_arrow_position()

    def set_error_state(self, message):
        self.co2_value_label.setText("----")
        self.status_text_label.setText(message)
        self.status_text_label.setStyleSheet(
            "background-color:#FFCDD2; border-radius:8px; color:#B71C1C;"
        )
        self.current_value = self.LEVELS[0]["min"]
        self.arrow_label.hide()
        self.update_arrow_position()

    def update_arrow_position(self):
        if not self.level_bars or self.arrow_container.width() <= 0:
            return

        value = self.current_value
        index = self.get_level_index(value)
        level = self.LEVELS[index]

        level_min = level["min"]
        level_max = level["max"]

        if value <= level_min:
            ratio = 0.0
        elif value >= level_max:
            ratio = 1.0
        else:
            denominator = level_max - level_min
            ratio = 0.0 if denominator <= 0 else (value - level_min) / denominator

        target_bar = self.level_bars[index]
        bar_x = target_bar.geometry().x()
        bar_width = target_bar.geometry().width()

        target_x = bar_x + bar_width * ratio
        arrow_x = int(target_x - self.arrow_label.width() / 2)
        max_x = self.arrow_container.width() - self.arrow_label.width()
        arrow_x = max(0, min(arrow_x, max_x))

        arrow_y = self.arrow_container.height() - self.arrow_label.height()
        self.arrow_label.move(QPoint(arrow_x, int(arrow_y)))

    def resizeEvent(self, event):
        super().resizeEvent(event)
        # [최적화 2] 50ms 딜레이를 주어 리사이즈 중 과도한 연산 방지
        self.resize_timer.start(50)


# ============================================================
# 9. 메인 모니터링 창
# ============================================================

class CO2MonitorApp(QMainWindow):

    def __init__(self, target_ports, db_manager):
        super().__init__()

        self.target_ports = target_ports
        self.db_manager = db_manager

        self.threads = []
        self.widgets = []
        self.dragPos = QPoint()
        self.is_always_on_top = False
        self.is_closing = False

        self.initUI()
        self.start_threads()

    def initUI(self):
        self.setWindowTitle("이산화탄소 다중 모니터링")
        num_ports = len(self.target_ports)
        window_width = max(280, num_ports * 300)

        self.resize(window_width, 380)
        self.setMinimumSize(280, 300)

        self.setStyleSheet("""
            QMainWindow {
                background-color: #E9ECEF;
                border: none;
            }
        """)

        self.setWindowFlags(Qt.FramelessWindowHint)

        central_widget = QWidget()
        self.setCentralWidget(central_widget)

        main_layout = QHBoxLayout(central_widget)
        main_layout.setContentsMargins(15, 15, 15, 15)
        main_layout.setSpacing(15)

        for i, port in enumerate(self.target_ports):
            widget = CO2LevelWidget(
                sensor_index=i,
                title=f"센서 {i + 1} ({port})"
            )
            self.widgets.append(widget)
            main_layout.addWidget(widget)

    def start_threads(self):
        for i, port in enumerate(self.target_ports):
            thread = SerialThread(
                port_name=port,
                sensor_index=i,
                db_manager=self.db_manager
            )
            thread.data_signal.connect(self.update_data)
            thread.error_signal.connect(self.handle_error)
            self.threads.append(thread)

        for thread in self.threads:
            thread.start()

    def update_data(self, sensor_index, co2_value):
        if 0 <= sensor_index < len(self.widgets):
            self.widgets[sensor_index].update_co2(co2_value)

    def handle_error(self, sensor_index, error_message):
        if 0 <= sensor_index < len(self.widgets):
            self.widgets[sensor_index].set_error_state(error_message)

    def contextMenuEvent(self, event):
        menu = QMenu(self)

        top_action = QAction("최상단 고정 켜기/끄기", self)
        top_action.triggered.connect(self.toggle_always_on_top)
        menu.addAction(top_action)

        menu.addSeparator()

        capture_action = QAction("화면 캡처", self)
        capture_action.triggered.connect(self.capture_screen)
        menu.addAction(capture_action)

        menu.addSeparator()

        exit_action = QAction("종료", self)
        exit_action.triggered.connect(self.close)
        menu.addAction(exit_action)

        menu.exec_(event.globalPos())

    def toggle_always_on_top(self):
        self.is_always_on_top = not self.is_always_on_top
        flags = self.windowFlags()

        if self.is_always_on_top:
            flags |= Qt.WindowStaysOnTopHint
        else:
            flags &= ~Qt.WindowStaysOnTopHint

        self.setWindowFlags(flags)
        self.show()

    def capture_screen(self):
        try:
            screen = QApplication.primaryScreen()
            if screen is None:
                return

            screenshot = screen.grabWindow(self.winId())
            save_dir = os.path.join(BASE_DIR, "Captures")
            os.makedirs(save_dir, exist_ok=True)

            filename = datetime.now().strftime("capture_%Y%m%d_%H%M%S.png")
            file_path = os.path.join(save_dir, filename)

            screenshot.save(file_path, "PNG")

            QMessageBox.information(
                self,
                "캡처 완료",
                f"화면이 저장되었습니다.\n\n{file_path}"
            )
        except Exception as e:
            QMessageBox.warning(self, "캡처 오류", str(e))

    def closeEvent(self, event):
        if self.is_closing:
            event.accept()
            return

        self.is_closing = True

        for thread in self.threads:
            if thread.isRunning():
                thread.requestInterruption()

        for thread in self.threads:
            if thread.isRunning():
                thread.wait(3000)

        try:
            self.db_manager.export_to_excel()
        except Exception as e:
            print(f"Excel 내보내기 실패: {e}")

        event.accept()

    def mousePressEvent(self, event):
        if event.button() == Qt.LeftButton:
            self.dragPos = event.globalPos() - self.frameGeometry().topLeft()
        super().mousePressEvent(event)

    def mouseMoveEvent(self, event):
        if event.buttons() & Qt.LeftButton:
            self.move(event.globalPos() - self.dragPos)
        super().mouseMoveEvent(event)

    def keyPressEvent(self, event):
        if event.key() == Qt.Key_Escape:
            self.close()
        else:
            super().keyPressEvent(event)


# ============================================================
# 10. 프로그램 시작
# ============================================================

def main():
    app = QApplication(sys.argv)

    db_manager = DatabaseManager()
    db_manager.cleanup_old_data()
    FileLogger.cleanup_old_files()

    setup_dialog = PortSelectionDialog()
    result = setup_dialog.exec_()

    if result != QDialog.Accepted:
        sys.exit(0)

    selected_ports = setup_dialog.get_selected_ports()

    if not selected_ports:
        QMessageBox.warning(None, "오류", "선택된 포트가 없습니다.")
        sys.exit(0)

    window = CO2MonitorApp(selected_ports, db_manager)
    window.show()

    sys.exit(app.exec_())


if __name__ == "__main__":
    main()