import os
import sys
import time
import sqlite3
import traceback
import queue
import threading
import json
import math
import serial
import serial.tools.list_ports

APP_VERSION = "4.0.0"

from collections import deque
from datetime import datetime, timedelta

from PyQt5.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QLabel, QFrame, QPushButton, QDialog, QMessageBox, QScrollArea,
    QGroupBox, QComboBox, QLineEdit, QProgressDialog
)
from PyQt5.QtCore import QThread, pyqtSignal, Qt, QPoint, QTimer
from PyQt5.QtGui import QFont

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# ============================================================
# Config
# ============================================================

class Config:
    BAUD_RATE = 9600
    SERIAL_TIMEOUT = 0.1
    MAX_SENSORS = 4
    SMOOTHING_WINDOW = 1
    NO_DATA_TIMEOUT = 30.0
    RECONNECT_DELAY_MS = 3000
    LOG_RETENTION_DAYS = 30
    THREAD_WAIT_MS = 5000
    DB_QUEUE_SIZE = 2000
    DB_QUEUE_PUT_TIMEOUT = 2.0
    DB_RETRY_COUNT = 3
    DB_RETRY_DELAY = 0.5
    SERIAL_ERROR_RECONNECT_DELAY_MS = 1000

    DATA_DIR = os.path.join(BASE_DIR, "Data")
    LOG_DIR = os.path.join(BASE_DIR, "Logs")
    EXCEL_DIR = os.path.join(BASE_DIR, "Excel_Logs")
    CALIBRATION_FILE = os.path.join(DATA_DIR, "sensor_calibration.json")
    VACUUM_MIN_INTERVAL_DAYS = 7

    # 🟢 [핵심] 기능검사 ±30% 통과율 95%를 달성한 센서별 최적 보정값 기본 탑재
    DEFAULT_SENSOR_CALIBRATION = {
        0: {"threshold": (3.5, 3.5, 3.5), "low_scale": (0.76, 0.76, 0.76), "low_offset": (6.18, 6.18, 6.18), "high_scale": (0.15, 0.15, 0.15), "high_offset": (9.21, 9.21, 9.21)},
        1: {"threshold": (3.5, 3.5, 3.5), "low_scale": (0.79, 0.79, 0.79), "low_offset": (6.12, 6.12, 6.12), "high_scale": (0.15, 0.15, 0.15), "high_offset": (9.20, 9.20, 9.20)},
        2: {"threshold": (4.0, 4.0, 4.0), "low_scale": (0.61, 0.61, 0.61), "low_offset": (6.10, 6.10, 6.10), "high_scale": (0.15, 0.15, 0.15), "high_offset": (9.17, 9.17, 9.17)},
        3: {"threshold": (4.5, 4.5, 4.5), "low_scale": (0.68, 0.68, 0.68), "low_offset": (6.03, 6.03, 6.03), "high_scale": (0.15, 0.15, 0.15), "high_offset": (9.18, 9.18, 9.18)},
    }
    SENSOR_CALIBRATION = {}

    PM10_LEVELS = [
        {"name": "좋음", "min": 0, "max": 30, "color": "#28A745"},
        {"name": "보통", "min": 31, "max": 80, "color": "#FFD700"},
        {"name": "민감군", "min": 81, "max": 120, "color": "#FD7E14"},
        {"name": "나쁨", "min": 121, "max": 150, "color": "#DC3545"},
        {"name": "매우 나쁨", "min": 151, "max": 300, "color": "#800080"},
        {"name": "위험", "min": 301, "max": 600, "color": "#795548"},
    ]

    PM25_LEVELS = [
        {"name": "좋음", "min": 0, "max": 15, "color": "#28A745"},
        {"name": "보통", "min": 16, "max": 35, "color": "#FFD700"},
        {"name": "민감군", "min": 36, "max": 50, "color": "#FD7E14"},
        {"name": "나쁨", "min": 51, "max": 75, "color": "#DC3545"},
        {"name": "매우 나쁨", "min": 76, "max": 100, "color": "#800080"},
        {"name": "위험", "min": 101, "max": 500, "color": "#795548"},
    ]

    PM1_LEVELS = [
        {"name": "좋음", "min": 0, "max": 10, "color": "#28A745"},
        {"name": "보통", "min": 11, "max": 25, "color": "#FFD700"},
        {"name": "민감군", "min": 26, "max": 35, "color": "#FD7E14"},
        {"name": "나쁨", "min": 36, "max": 50, "color": "#DC3545"},
        {"name": "매우 나쁨", "min": 51, "max": 75, "color": "#800080"},
        {"name": "위험", "min": 76, "max": 300, "color": "#795548"},
    ]

    @classmethod
    def _normalize_calibration(cls, data):
        normalized = {}
        defaults = cls.DEFAULT_SENSOR_CALIBRATION
        for i in range(cls.MAX_SENSORS):
            if isinstance(data, dict):
                raw = data.get(str(i), data.get(i, defaults.get(i)))
            else:
                raw = defaults.get(i)
            try:
                # 하위 호환 및 무결성 검증
                th = tuple(float(x) for x in raw.get("threshold", defaults[i]["threshold"])[:3])
                ls = tuple(float(x) for x in raw.get("low_scale", defaults[i]["low_scale"])[:3])
                lo = tuple(float(x) for x in raw.get("low_offset", defaults[i]["low_offset"])[:3])
                hs = tuple(float(x) for x in raw.get("high_scale", defaults[i]["high_scale"])[:3])
                ho = tuple(float(x) for x in raw.get("high_offset", defaults[i]["high_offset"])[:3])

                if len(th) != 3 or len(ls) != 3 or len(lo) != 3 or len(hs) != 3 or len(ho) != 3:
                    raise ValueError
                if not all(math.isfinite(x) for x in th + ls + lo + hs + ho):
                    raise ValueError
                if any(x < 0 for x in ls + hs):
                    raise ValueError
                normalized[i] = {"threshold": th, "low_scale": ls, "low_offset": lo, "high_scale": hs, "high_offset": ho}
            except (KeyError, TypeError, ValueError, IndexError):
                normalized[i] = {
                    "threshold": tuple(defaults[i]["threshold"]),
                    "low_scale": tuple(defaults[i]["low_scale"]),
                    "low_offset": tuple(defaults[i]["low_offset"]),
                    "high_scale": tuple(defaults[i]["high_scale"]),
                    "high_offset": tuple(defaults[i]["high_offset"])
                }
        return normalized

    @classmethod
    def load_calibration(cls):
        os.makedirs(cls.DATA_DIR, exist_ok=True)
        try:
            if os.path.exists(cls.CALIBRATION_FILE):
                with open(cls.CALIBRATION_FILE, "r", encoding="utf-8") as f:
                    data = json.load(f)
                cls.SENSOR_CALIBRATION = cls._normalize_calibration(data)
                return
        except (OSError, json.JSONDecodeError) as e:
            print(f"보정값 파일 읽기 오류: {e}")
        cls.SENSOR_CALIBRATION = cls._normalize_calibration(cls.DEFAULT_SENSOR_CALIBRATION)
        cls.save_calibration()

    @classmethod
    def save_calibration(cls):
        os.makedirs(cls.DATA_DIR, exist_ok=True)
        serializable = {
            str(i): {
                "threshold": list(params["threshold"]),
                "low_scale": list(params["low_scale"]),
                "low_offset": list(params["low_offset"]),
                "high_scale": list(params["high_scale"]),
                "high_offset": list(params["high_offset"])
            } for i, params in cls.SENSOR_CALIBRATION.items()
        }
        temp_file = cls.CALIBRATION_FILE + ".tmp"
        try:
            with open(temp_file, "w", encoding="utf-8") as f:
                json.dump(serializable, f, ensure_ascii=False, indent=2)
            os.replace(temp_file, cls.CALIBRATION_FILE)
            return True
        except OSError as e:
            print(f"보정값 파일 저장 오류: {e}")
            return False


# ============================================================
# Dust Parser
# ============================================================

class DustParser:
    @staticmethod
    def parse(raw_data, calib_params):
        try:
            raw_data = raw_data.strip()
            if not raw_data:
                return None

            parts = [item.strip() for item in raw_data.split(",")]
            if len(parts) < 3:
                return None

            thresholds = calib_params.get("threshold", (4.0, 4.0, 4.0))
            low_scales = calib_params.get("low_scale", (1.0, 1.0, 1.0))
            low_offsets = calib_params.get("low_offset", (0.0, 0.0, 0.0))
            high_scales = calib_params.get("high_scale", (1.0, 1.0, 1.0))
            high_offsets = calib_params.get("high_offset", (0.0, 0.0, 0.0))

            raw_pm1 = float(parts[0])
            raw_pm25 = float(parts[1])
            raw_pm10 = float(parts[2])

            raw_values = (raw_pm10, raw_pm25, raw_pm1)
            if not all(math.isfinite(v) for v in raw_values):
                return None

            # 저/고농도 분리 보정 함수
            def apply_piecewise_calibration(raw_val, th, l_scale, l_offset, h_scale, h_offset):
                if raw_val < th:
                    return (raw_val * l_scale) + l_offset
                else:
                    return (raw_val * h_scale) + h_offset

            calibrated_pm10 = apply_piecewise_calibration(raw_pm10, thresholds[0], low_scales[0], low_offsets[0], high_scales[0], high_offsets[0])
            calibrated_pm25 = apply_piecewise_calibration(raw_pm25, thresholds[1], low_scales[1], low_offsets[1], high_scales[1], high_offsets[1])
            calibrated_pm1 = apply_piecewise_calibration(raw_pm1, thresholds[2], low_scales[2], low_offsets[2], high_scales[2], high_offsets[2])

            calibrated_values = (calibrated_pm10, calibrated_pm25, calibrated_pm1)
            if not all(math.isfinite(v) for v in calibrated_values):
                return None

            pm10 = max(0, int(round(calibrated_pm10)))
            pm25 = max(0, int(round(calibrated_pm25)))
            pm1 = max(0, int(round(calibrated_pm1)))

            if pm1 > 1000 or pm25 > 1000 or pm10 > 2000:
                return "OUT_OF_RANGE"

            return {
                "pm10": pm10,
                "pm25": pm25,
                "pm1": pm1,
                "raw_pm10": raw_pm10,
                "raw_pm25": raw_pm25,
                "raw_pm1": raw_pm1,
            }
        except (ValueError, TypeError, IndexError):
            return None


# ============================================================
# Database Manager
# ============================================================

class DatabaseManager:
    _instance = None
    _lock = threading.Lock()

    def __new__(cls, db_path=None):
        with cls._lock:
            if cls._instance is None:
                cls._instance = super(DatabaseManager, cls).__new__(cls)
                cls._instance._initialized = False
        return cls._instance

    def __init__(self, db_path=None):
        if self._initialized:
            return
        with self._lock:
            if self._initialized:
                return
            self.db_path = db_path or os.path.join(BASE_DIR, "Data", "dust_measurement.db")
            self.queue = queue.Queue(maxsize=Config.DB_QUEUE_SIZE)
            self.running = True
            self._stop_lock = threading.Lock()
            self._stopped = False
            self._init_db()
            self.worker_thread = threading.Thread(target=self._db_worker, name="SQLiteWorker", daemon=True)
            self.worker_thread.start()
            self._initialized = True

    def _connect(self):
        conn = sqlite3.connect(self.db_path, timeout=15.0)
        conn.execute("PRAGMA busy_timeout=10000;")
        return conn

    def _init_db(self):
        try:
            os.makedirs(os.path.dirname(self.db_path), exist_ok=True)
            with sqlite3.connect(self.db_path, timeout=15.0) as conn:
                conn.execute("PRAGMA journal_mode=WAL;")
                conn.execute("PRAGMA synchronous=NORMAL;")
                conn.execute("""
                    CREATE TABLE IF NOT EXISTS measurements (
                        id INTEGER PRIMARY KEY AUTOINCREMENT,
                        measured_at TEXT NOT NULL,
                        sensor_index INTEGER NOT NULL,
                        port TEXT NOT NULL,
                        pm10 REAL,
                        pm25 REAL,
                        pm1 REAL,
                        raw_pm10 REAL,
                        raw_pm25 REAL,
                        raw_pm1 REAL,
                        status TEXT DEFAULT 'NORMAL'
                    )
                """)

                existing_columns = {row[1] for row in conn.execute("PRAGMA table_info(measurements)").fetchall()}
                required_columns = {
                    "raw_pm10": "REAL",
                    "raw_pm25": "REAL",
                    "raw_pm1": "REAL",
                    "status": "TEXT DEFAULT 'NORMAL'",
                }
                for column_name, column_type in required_columns.items():
                    if column_name not in existing_columns:
                        conn.execute(f"ALTER TABLE measurements ADD COLUMN {column_name} {column_type}")
                
                conn.execute("CREATE INDEX IF NOT EXISTS idx_measurements_time ON measurements(measured_at)")
                conn.execute("CREATE INDEX IF NOT EXISTS idx_measurements_sensor ON measurements(sensor_index)")
                conn.execute("CREATE INDEX IF NOT EXISTS idx_measurements_sensor_time ON measurements(sensor_index, measured_at)")
                conn.commit()
        except Exception as e:
            print(f"Global SQLite 초기화 오류: {e}")

    def _db_worker(self):
        conn = None
        while self.running or not self.queue.empty():
            try:
                if conn is None:
                    try:
                        conn = self._connect()
                    except Exception as e:
                        print(f"SQLite 연결 실패: {type(e).__name__}: {e}")
                        if not self.running:
                            break
                        time.sleep(1.0)
                        continue

                try:
                    task = self.queue.get(timeout=0.5)
                except queue.Empty:
                    continue

                if task is None:
                    self.queue.task_done()
                    continue

                func, args, future_event, result_holder = task
                try:
                    result = func(conn, *args)
                    conn.commit()
                    if result_holder is not None:
                        result_holder["result"] = result
                except Exception as e:
                    try:
                        conn.rollback()
                    except Exception:
                        pass
                    if result_holder is not None:
                        result_holder["error"] = e
                    print(f"SQLite 작업 오류: {type(e).__name__}: {e}")
                    try:
                        conn.close()
                    except Exception:
                        pass
                    conn = None
                finally:
                    if future_event is not None:
                        future_event.set()
                    self.queue.task_done()
            except Exception as e:
                print(f"DB Worker 루프 오류: {type(e).__name__}: {e}")
                if conn is not None:
                    try:
                        conn.close()
                    except Exception:
                        pass
                    conn = None
                time.sleep(0.2)

        if conn is not None:
            try:
                conn.close()
            except Exception:
                pass

    def is_worker_alive(self):
        return self.worker_thread is not None and self.worker_thread.is_alive()

    def execute_async(self, func, *args):
        with self._stop_lock:
            if self._stopped or not self.running:
                return False
            task = (func, args, None, None)
            try:
                self.queue.put(task, timeout=Config.DB_QUEUE_PUT_TIMEOUT)
                return True
            except queue.Full:
                print("SQLite queue가 가득 차서 비동기 작업을 넣지 못했습니다.")
                return False

    def execute_sync(self, func, *args, timeout=5.0):
        future_event = threading.Event()
        result_holder = {}
        with self._stop_lock:
            if self._stopped or not self.running:
                return None
            try:
                self.queue.put(
                    (func, args, future_event, result_holder),
                    timeout=Config.DB_QUEUE_PUT_TIMEOUT
                )
            except queue.Full:
                print("SQLite queue가 가득 차서 동기 작업을 넣지 못했습니다.")
                return None

        if not future_event.wait(timeout=timeout):
            print("SQLite 작업 대기 시간이 초과되었습니다.")
            return None
        if "error" in result_holder:
            print(f"SQLite 동기 작업 실패: {result_holder['error']}")
            return None
        return result_holder.get("result")

    def flush(self, timeout=10.0):
        done = threading.Event()
        result_holder = {}

        def barrier(conn):
            return True

        with self._stop_lock:
            if self._stopped:
                return False
            try:
                self.queue.put(
                    (barrier, (), done, result_holder),
                    timeout=Config.DB_QUEUE_PUT_TIMEOUT
                )
            except queue.Full:
                return False

        if not done.wait(timeout=timeout):
            return False
        return "error" not in result_holder

    def stop(self, timeout=10.0):
        with self._stop_lock:
            if self._stopped:
                return True
            self.running = False

        try:
            self.queue.put(None, timeout=Config.DB_QUEUE_PUT_TIMEOUT)
        except queue.Full:
            pass

        self.worker_thread.join(timeout=timeout)
        alive = self.worker_thread.is_alive()
        with self._stop_lock:
            self._stopped = True
        return not alive


# ============================================================
# Sensor Logger
# ============================================================

class SensorLogger:
    _vacuum_lock = threading.Lock()

    def __init__(self, port_name, sensor_index):
        self.port_name = port_name
        self.sensor_index = sensor_index
        self.log_dir = Config.LOG_DIR
        self.excel_dir = Config.EXCEL_DIR
        self.db_dir = Config.DATA_DIR
        os.makedirs(self.log_dir, exist_ok=True)
        os.makedirs(self.excel_dir, exist_ok=True)
        os.makedirs(self.db_dir, exist_ok=True)
        self.db_path = os.path.join(self.db_dir, "dust_measurement.db")
        self.db_manager = DatabaseManager(self.db_path)

    @staticmethod
    def init_global_database(db_path):
        DatabaseManager(db_path)

    @staticmethod
    def cleanup_old_logs_global():
        threshold_time = time.time() - (Config.LOG_RETENTION_DAYS * 86400)
        for target_dir in [Config.LOG_DIR, Config.EXCEL_DIR]:
            if not os.path.exists(target_dir): continue
            for filename in os.listdir(target_dir):
                file_path = os.path.join(target_dir, filename)
                try:
                    if os.path.isfile(file_path) and os.path.getmtime(file_path) < threshold_time:
                        os.remove(file_path)
                except OSError: pass

    def write_error(self, message):
        log_file = os.path.join(self.log_dir, f"error_log_Sensor{self.sensor_index + 1}.txt")
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        try:
            with open(log_file, "a", encoding="utf-8-sig") as f:
                f.write(f"{timestamp} | 포트: {self.port_name} | {message}\n")
        except OSError: pass

    def save_measurements_batch(self, measurements):
        if not measurements:
            return True
        data = list(measurements)

        def insert_rows(conn, rows):
            conn.executemany("""
                INSERT INTO measurements (measured_at, sensor_index, port, pm10, pm25, pm1, raw_pm10, raw_pm25, raw_pm1, status)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, rows)
            return True

        for attempt in range(Config.DB_RETRY_COUNT):
            if self.db_manager.execute_async(insert_rows, data):
                return True
            if attempt + 1 < Config.DB_RETRY_COUNT:
                time.sleep(Config.DB_RETRY_DELAY)
        self.write_error(f"DB 비동기 저장 실패: {len(data)}건")
        return False

    def save_measurement_sync(self, measurement, timeout=5.0):
        def insert_one(conn, row):
            conn.execute("""
                INSERT INTO measurements (measured_at, sensor_index, port, pm10, pm25, pm1, raw_pm10, raw_pm25, raw_pm1, status)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, row)
            return True
        return bool(self.db_manager.execute_sync(insert_one, measurement, timeout=timeout))

    def _fetch_rows_for_date(self, date_str):
        start_dt = f"{date_str} 00:00:00"
        next_dt = (datetime.strptime(date_str, "%Y-%m-%d") + timedelta(days=1)).strftime("%Y-%m-%d") + " 00:00:00"
        def fetch_rows(conn):
            cursor = conn.execute("""
                SELECT measured_at, port, pm10, pm25, pm1, raw_pm10, raw_pm25, raw_pm1, status
                FROM measurements
                WHERE sensor_index = ? AND measured_at >= ? AND measured_at < ?
                ORDER BY measured_at
            """, (self.sensor_index, start_dt, next_dt))
            return cursor.fetchall()
        return self.db_manager.execute_sync(fetch_rows, timeout=5.0)

    def _write_excel(self, date_str, rows):
        if not rows: return False
        filename = f"Dust_log_{date_str}_Sensor{self.sensor_index + 1}.xlsx"
        file_path = os.path.join(self.excel_dir, filename)
        temp_path = file_path + ".tmp"

        wb = Workbook()
        ws = wb.active
        ws.title = f"{date_str} 측정 데이터"

        headers = ["측정일시", "포트", "PM10", "PM2.5", "PM1.0", "상태"]
        ws.append(headers)

        thin_border = Border(left=Side(style="thin", color="D3D3D3"), right=Side(style="thin", color="D3D3D3"),
                             top=Side(style="thin", color="D3D3D3"), bottom=Side(style="thin", color="D3D3D3"))
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(name="Malgun Gothic", size=10, bold=True, color="FFFFFF")
        data_font = Font(name="Malgun Gothic", size=9)

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        for row_data in rows:
            # 원본(raw) 데이터 제외하고 6개만 엑셀에 출력
            filtered_row = [row_data[0], row_data[1], row_data[2], row_data[3], row_data[4], row_data[8]]
            ws.append(filtered_row)
            
            current_row = ws.max_row
            for col_num, value in enumerate(filtered_row, 1):
                cell = ws.cell(row=current_row, column=col_num)
                cell.font = data_font
                cell.border = thin_border
                if col_num in (1, 2, 6): 
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else: 
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    if isinstance(value, (int, float)): 
                        cell.number_format = "#,##0"

        for i, width in enumerate([19, 14, 10, 10, 10, 12], 1):
            ws.column_dimensions[get_column_letter(i)].width = width

        try:
            wb.save(temp_path)
            os.replace(temp_path, file_path)
            return True
        finally:
            try:
                if os.path.exists(temp_path): os.remove(temp_path)
            except OSError: pass

    def export_excel(self, date_str=None):
        date_str = date_str or datetime.now().strftime("%Y-%m-%d")
        rows = self._fetch_rows_for_date(date_str)
        return self._write_excel(date_str, rows) if rows else False


# ============================================================
# Serial Thread
# ============================================================

class SerialThread(QThread):
    data_signal = pyqtSignal(int, int, int, int)
    error_signal = pyqtSignal(int, str)
    date_changed_signal = pyqtSignal(int, str)

    def __init__(self, port_name, sensor_index):
        super().__init__()
        self.port_name = port_name
        self.sensor_index = sensor_index
        
        default_calib = Config.DEFAULT_SENSOR_CALIBRATION.get(sensor_index)
        stored = Config.SENSOR_CALIBRATION.get(sensor_index, default_calib)
        self.calib_params = {
            "threshold": tuple(stored.get("threshold", default_calib["threshold"])),
            "low_scale": tuple(stored.get("low_scale", default_calib["low_scale"])),
            "low_offset": tuple(stored.get("low_offset", default_calib["low_offset"])),
            "high_scale": tuple(stored.get("high_scale", default_calib["high_scale"])),
            "high_offset": tuple(stored.get("high_offset", default_calib["high_offset"]))
        }
        
        self.logger = SensorLogger(port_name, sensor_index)
        window_size = max(1, Config.SMOOTHING_WINDOW)

        self.data_buffer = {
            "pm10": deque(maxlen=window_size),
            "pm25": deque(maxlen=window_size),
            "pm1": deque(maxlen=window_size)
        }
        
        self.minute_sum = {"pm10": 0, "pm25": 0, "pm1": 0, "raw_pm10": 0.0, "raw_pm25": 0.0, "raw_pm1": 0.0}
        self.minute_count = 0
        self.current_minute_bucket = datetime.now().replace(second=0, microsecond=0)
        self._buffer_lock = threading.Lock()
        self.last_valid_data_time = time.time()
        self.db_save_failures = 0

    def safe_sleep(self, milliseconds):
        remaining = milliseconds
        while remaining > 0:
            if self.isInterruptionRequested(): return
            sleep_time = min(50, remaining)
            self.msleep(sleep_time)
            remaining -= sleep_time

    def clear_ui_buffers(self):
        with self._buffer_lock:
            for buffer in self.data_buffer.values(): buffer.clear()

    def clear_all_buffers(self):
        with self._buffer_lock:
            for buffer in self.data_buffer.values(): buffer.clear()
            self.minute_sum = {"pm10": 0, "pm25": 0, "pm1": 0, "raw_pm10": 0.0, "raw_pm25": 0.0, "raw_pm1": 0.0}
            self.minute_count = 0

    def add_measurement(self, pm10, pm25, pm1, raw_pm10, raw_pm25, raw_pm1):
        with self._buffer_lock:
            self.data_buffer["pm10"].append(pm10)
            self.data_buffer["pm25"].append(pm25)
            self.data_buffer["pm1"].append(pm1)

            self.minute_sum["pm10"] += pm10
            self.minute_sum["pm25"] += pm25
            self.minute_sum["pm1"] += pm1
            self.minute_sum["raw_pm10"] += raw_pm10
            self.minute_sum["raw_pm25"] += raw_pm25
            self.minute_sum["raw_pm1"] += raw_pm1
            self.minute_count += 1

    def save_current_minute_average(self, minute_bucket, synchronous=False):
        with self._buffer_lock:
            count = self.minute_count
            if count <= 0: return True
            avg_pm10 = int(round(self.minute_sum["pm10"] / count))
            avg_pm25 = int(round(self.minute_sum["pm25"] / count))
            avg_pm1 = int(round(self.minute_sum["pm1"] / count))
            
            avg_raw_pm10 = round(self.minute_sum["raw_pm10"] / count, 2)
            avg_raw_pm25 = round(self.minute_sum["raw_pm25"] / count, 2)
            avg_raw_pm1 = round(self.minute_sum["raw_pm1"] / count, 2)

        measured_at = minute_bucket.strftime("%Y-%m-%d %H:%M:00")
        row = (measured_at, self.sensor_index, self.port_name, 
               avg_pm10, avg_pm25, avg_pm1, 
               avg_raw_pm10, avg_raw_pm25, avg_raw_pm1, 
               "NORMAL")

        success = self.logger.save_measurement_sync(row, timeout=5.0) if synchronous else self.logger.save_measurements_batch([row])
        if success:
            self.db_save_failures = 0
            with self._buffer_lock:
                if self.minute_count == count:
                    self.minute_sum = {"pm10": 0, "pm25": 0, "pm1": 0, "raw_pm10": 0.0, "raw_pm25": 0.0, "raw_pm1": 0.0}
                    self.minute_count = 0
            return True
        self.db_save_failures += 1
        self.logger.write_error(
            f"분 평균 DB 저장 실패: bucket={minute_bucket.strftime('%Y-%m-%d %H:%M:%S')}, failures={self.db_save_failures}"
        )
        return False

    def check_minute_boundary(self):
        new_bucket = datetime.now().replace(second=0, microsecond=0)
        if new_bucket == self.current_minute_bucket: return
        old_bucket = self.current_minute_bucket
        if self.minute_count > 0:
            if not self.save_current_minute_average(old_bucket, synchronous=False): return
        if old_bucket.date() != new_bucket.date():
            self.date_changed_signal.emit(self.sensor_index, old_bucket.strftime("%Y-%m-%d"))
        self.current_minute_bucket = new_bucket

    def get_ui_averages(self):
        with self._buffer_lock:
            pm10, pm25, pm1 = list(self.data_buffer["pm10"]), list(self.data_buffer["pm25"]), list(self.data_buffer["pm1"])
        if not pm10 or not pm25 or not pm1: return None
        return (int(sum(pm10) / len(pm10)), int(sum(pm25) / len(pm25)), int(sum(pm1) / len(pm1)))

    def run(self):
        is_connected = False
        no_data_error_sent = False
        last_ui_update_time = 0.0
        last_data_time = time.time()

        while not self.isInterruptionRequested():
            ser = None
            try:
                ser = serial.Serial(self.port_name, Config.BAUD_RATE, timeout=Config.SERIAL_TIMEOUT)
                ser.reset_input_buffer()
                ser.readline()
                is_connected = True
                no_data_error_sent = False
                last_data_time = time.time()
                self.current_minute_bucket = datetime.now().replace(second=0, microsecond=0)
                self.clear_ui_buffers()

                while ser.is_open and not self.isInterruptionRequested():
                    parsed_values = []
                    while ser.in_waiting > 0:
                        if self.isInterruptionRequested(): break
                        raw_data = ser.readline().decode("utf-8", errors="ignore").strip()
                        if not raw_data: continue
                        parsed = DustParser.parse(raw_data, self.calib_params)
                        if parsed == "OUT_OF_RANGE": continue
                        if parsed is not None:
                            parsed_values.append(parsed)
                            last_data_time = time.time()

                    self.check_minute_boundary()

                    if time.time() - last_data_time >= Config.NO_DATA_TIMEOUT:
                        if not no_data_error_sent:
                            self.error_signal.emit(self.sensor_index, "데이터 없음")
                            no_data_error_sent = True
                            self.clear_ui_buffers()
                    elif parsed_values:
                        no_data_error_sent = False
                        for p in parsed_values:
                            self.add_measurement(p["pm10"], p["pm25"], p["pm1"], p["raw_pm10"], p["raw_pm25"], p["raw_pm1"])
                        
                        current_time = time.time()
                        if current_time - last_ui_update_time >= 1.0:
                            last_ui_update_time = current_time
                            avgs = self.get_ui_averages()
                            if avgs: self.data_signal.emit(self.sensor_index, avgs[0], avgs[1], avgs[2])
                    self.safe_sleep(100)
            except serial.SerialException as e:
                is_connected = False
                self.clear_all_buffers()
                self.logger.write_error(f"Serial 오류: {type(e).__name__}: {e}")
                self.error_signal.emit(self.sensor_index, "연결 실패")
                if not self.isInterruptionRequested():
                    self.safe_sleep(Config.SERIAL_ERROR_RECONNECT_DELAY_MS)
            except Exception as e:
                is_connected = False
                self.clear_all_buffers()
                self.logger.write_error(f"예상하지 못한 센서 스레드 오류: {traceback.format_exc()}")
                self.error_signal.emit(self.sensor_index, "센서 오류")
            finally:
                if ser is not None:
                    try:
                        ser.close()
                    except Exception:
                        pass
                if not self.isInterruptionRequested():
                    self.safe_sleep(Config.RECONNECT_DELAY_MS)
        try: self.save_current_minute_average(self.current_minute_bucket, synchronous=True)
        except Exception: pass

    def update_calibration(self, new_calib_params):
        try:
            th = tuple(float(x) for x in new_calib_params["threshold"])
            ls = tuple(float(x) for x in new_calib_params["low_scale"])
            lo = tuple(float(x) for x in new_calib_params["low_offset"])
            hs = tuple(float(x) for x in new_calib_params["high_scale"])
            ho = tuple(float(x) for x in new_calib_params["high_offset"])

            if len(th) != 3 or len(ls) != 3 or len(lo) != 3 or len(hs) != 3 or len(ho) != 3:
                return False
            if not all(math.isfinite(x) for x in th + ls + lo + hs + ho):
                return False
            if any(x < 0 for x in ls + hs):
                return False

            if self.minute_count > 0:
                bucket = self.current_minute_bucket
                if not self.save_current_minute_average(bucket, synchronous=True):
                    self.logger.write_error("보정값 변경 전 미저장 분 데이터 저장 실패")
                    return False

            with self._buffer_lock:
                self.calib_params = {
                    "threshold": th,
                    "low_scale": ls,
                    "low_offset": lo,
                    "high_scale": hs,
                    "high_offset": ho
                }
                for buffer in self.data_buffer.values():
                    buffer.clear()
            return True
        except Exception as e:
            self.logger.write_error(f"보정값 적용 오류: {type(e).__name__}: {e}")
            return False


# ============================================================
# Dust Level Widget
# ============================================================

class DustLevelWidget(QWidget):

    def __init__(self, title="미세먼지", levels=None, parent=None):
        super().__init__(parent)
        self.title = title
        self.levels = levels if levels else Config.PM10_LEVELS
        self.current_value = 0
        self.initUI()

    def initUI(self):
        self.setStyleSheet("background-color: transparent; border: none;")
        main_layout = QVBoxLayout(self)
        main_layout.setContentsMargins(10, 10, 10, 10)
        main_layout.setSpacing(3)

        self.title_label = QLabel(self.title)
        self.title_label.setFont(QFont("Malgun Gothic", 11, QFont.Bold))
        self.title_label.setAlignment(Qt.AlignCenter)
        self.title_label.setStyleSheet("color: black; border: none;")
        main_layout.addWidget(self.title_label)

        value_layout = QHBoxLayout()
        value_layout.setAlignment(Qt.AlignCenter)

        self.val_label = QLabel("----")
        self.val_label.setFont(QFont("Arial", 28, QFont.Bold))
        self.val_label.setStyleSheet("color: black; border: none;")
        value_layout.addWidget(self.val_label)

        unit_label = QLabel("μg/m³")
        unit_label.setFont(QFont("Arial", 10, QFont.Bold))
        unit_label.setStyleSheet("color: #666666; margin-bottom: 4px; border: none;")
        unit_label.setAlignment(Qt.AlignBottom)
        value_layout.addWidget(unit_label)

        main_layout.addLayout(value_layout)

        self.arrow_container = QWidget()
        self.arrow_container.setStyleSheet("border: none;")
        self.arrow_container.setFixedHeight(14)

        self.arrow_label = QLabel("▼")
        self.arrow_label.setFont(QFont("Arial", 9, QFont.Bold))
        self.arrow_label.setStyleSheet("color: black; border: none;")
        self.arrow_label.setAlignment(Qt.AlignCenter)
        self.arrow_label.setFixedWidth(14)
        self.arrow_label.setParent(self.arrow_container)

        main_layout.addWidget(self.arrow_container)

        level_bar_frame = QFrame()
        level_bar_frame.setStyleSheet("border: none;")
        self.level_bar_layout = QHBoxLayout(level_bar_frame)
        self.level_bar_layout.setContentsMargins(0, 0, 0, 0)
        self.level_bar_layout.setSpacing(2)

        self.level_bars = []
        for i, level in enumerate(self.levels):
            bar = QFrame()
            bar.setFixedHeight(10)
            border_radius = ""
            if i == 0:
                border_radius = "border-top-left-radius: 5px; border-bottom-left-radius: 5px;"
            elif i == len(self.levels) - 1:
                border_radius = "border-top-right-radius: 5px; border-bottom-right-radius: 5px;"

            bar.setStyleSheet(f"background-color: {level['color']}; {border_radius}")
            self.level_bars.append(bar)
            self.level_bar_layout.addWidget(bar)

        main_layout.addWidget(level_bar_frame)

        self.status_text_label = QLabel("대기 중...")
        self.status_text_label.setFont(QFont("Malgun Gothic", 11, QFont.Bold))
        self.status_text_label.setAlignment(Qt.AlignCenter)
        self.status_text_label.setFixedHeight(30)
        self.status_text_label.setStyleSheet("background-color: #E0E0E0; border-radius: 6px; color: gray;")

        main_layout.addWidget(self.status_text_label)

    def update_val(self, value):
        self.arrow_label.show()
        self.current_value = value
        self.val_label.setText(str(value))

        current_level = self.levels[self._find_level_index(value)]
        self.status_text_label.setText(current_level["name"])

        text_color = "black" if current_level["name"] == "보통" else "white"
        self.status_text_label.setStyleSheet(
            f"background-color: {current_level['color']}; border-radius: 6px; color: {text_color}; border: none;"
        )
        self.update_arrow_position()

    def _find_level_index(self, value):
        for i, level in enumerate(self.levels):
            if value <= level["max"]:
                return i
        return len(self.levels) - 1

    def set_error_state(self, msg):
        self.val_label.setText("----")
        self.status_text_label.setText(msg)
        self.status_text_label.setStyleSheet("background-color: #FFCDD2; border-radius: 6px; color: #B71C1C; border: none;")
        self.current_value = self.levels[0]["min"]
        self.update_arrow_position()
        self.arrow_label.hide()

    def update_arrow_position(self):
        if not self.level_bars or self.level_bars[0].geometry().width() == 0:
            return

        value = self.current_value
        level_index = self._find_level_index(value)
        current_level = self.levels[level_index]

        level_min, level_max = current_level["min"], current_level["max"]

        if value <= level_min:
            ratio = 0.0
        elif value >= level_max:
            ratio = 1.0
        else:
            denominator = level_max - level_min
            ratio = 0.0 if denominator <= 0 else (value - level_min) / denominator

        target_bar = self.level_bars[level_index]
        target_x = target_bar.geometry().x() + (target_bar.geometry().width() * ratio)
        arrow_x = int(target_x - (self.arrow_label.width() / 2))
        max_x = self.arrow_container.width() - self.arrow_label.width()

        arrow_x = max(0, min(arrow_x, max_x))
        self.arrow_label.move(
            QPoint(arrow_x, int(self.arrow_container.height() - self.arrow_label.height() + 2))
        )

    def showEvent(self, event):
        super().showEvent(event)
        QTimer.singleShot(50, self.update_arrow_position)

    def resizeEvent(self, event):
        super().resizeEvent(event)
        self.update_arrow_position()


# ============================================================
# Port Selection Dialog
# ============================================================

class PortSelectionDialog(QDialog):

    def __init__(self, parent=None):
        super().__init__(parent)
        self.selected_ports = {}
        self.available_ports = sorted([port.device for port in serial.tools.list_ports.comports()])
        self.initUI()

    def initUI(self):
        self.setWindowTitle("센서 포트 설정")
        self.resize(350, 250)
        self.setStyleSheet("background-color: white;")

        layout = QVBoxLayout(self)
        layout.setContentsMargins(20, 20, 20, 20)
        layout.setSpacing(15)

        title_label = QLabel("모니터링할 센서 포트를 선택하세요")
        title_label.setFont(QFont("Malgun Gothic", 10, QFont.Bold))
        layout.addWidget(title_label)

        self.combos = []
        for i in range(Config.MAX_SENSORS):
            h_layout = QHBoxLayout()
            label = QLabel(f"센서 {i + 1} 포트:")
            label.setFont(QFont("Malgun Gothic", 9))

            combo = QComboBox()
            combo.addItem("선택 안 함")
            combo.addItems(self.available_ports)
            combo.setFont(QFont("Malgun Gothic", 9))
            combo.currentIndexChanged.connect(self.update_combo_items)

            h_layout.addWidget(label)
            h_layout.addWidget(combo)
            layout.addLayout(h_layout)

            self.combos.append(combo)

        self.start_btn = QPushButton("모니터링 시작")
        self.start_btn.setFixedHeight(35)
        self.start_btn.setFont(QFont("Malgun Gothic", 10, QFont.Bold))
        self.start_btn.setStyleSheet("background-color: #007BFF; color: white; border-radius: 4px;")
        self.start_btn.clicked.connect(self.accept_selection)

        layout.addWidget(self.start_btn)

    def update_combo_items(self):
        selected_values = [
            combo.currentText()
            for combo in self.combos
            if combo.currentText() and combo.currentText() != "선택 안 함"
        ]

        for combo in self.combos:
            current_selection = combo.currentText()
            combo.blockSignals(True)
            combo.clear()
            combo.addItem("선택 안 함")

            for port in self.available_ports:
                if port in selected_values and port != current_selection:
                    continue
                combo.addItem(port)

            index = combo.findText(current_selection)
            combo.setCurrentIndex(index if index != -1 else 0)
            combo.blockSignals(False)

    def accept_selection(self):
        self.selected_ports = {
            i: combo.currentText()
            for i, combo in enumerate(self.combos)
            if combo.currentText() and combo.currentText() != "선택 안 함"
        }

        if not self.selected_ports:
            QMessageBox.warning(self, "경고", "최소 하나 이상의 센서 포트를 선택해 주세요.")
            return

        self.accept()


# ============================================================
# Calibration Dialog
# ============================================================

class CalibrationDialog(QDialog):

    def __init__(self, current_calib_dict, parent=None):
        super().__init__(parent)
        self.calib_dict = {
            i: {
                "threshold": tuple(v["threshold"]),
                "low_scale": tuple(v["low_scale"]),
                "low_offset": tuple(v["low_offset"]),
                "high_scale": tuple(v["high_scale"]),
                "high_offset": tuple(v["high_offset"])
            }
            for i, v in current_calib_dict.items()
        }
        self.initUI()

    def initUI(self):
        self.setWindowTitle("센서 구간별 보정값 설정")
        self.resize(750, 550)
        self.setStyleSheet("background-color: white;")

        main_layout = QVBoxLayout(self)
        main_layout.setContentsMargins(20, 20, 20, 20)
        main_layout.setSpacing(15)

        title_label = QLabel("센서별 저농도/고농도 분리 보정 계수 설정")
        title_label.setFont(QFont("Malgun Gothic", 10, QFont.Bold))
        main_layout.addWidget(title_label)

        formula_label = QLabel("📌 계산 공식:\n 원본값 < TH 이면, (원본값 × Low S) + Low O\n 원본값 >= TH 이면, (원본값 × High S) + High O")
        formula_label.setFont(QFont("Malgun Gothic", 9))
        formula_label.setStyleSheet("color: #0056b3; background-color: #e7f1ff; padding: 6px; border-radius: 4px;")
        main_layout.addWidget(formula_label)

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)

        scroll_content = QWidget()
        scroll_layout = QVBoxLayout(scroll_content)
        scroll_layout.setSpacing(10)

        self.inputs = {}
        for i in range(Config.MAX_SENSORS):
            group = QGroupBox(f"센서 {i + 1}")
            group.setFont(QFont("Malgun Gothic", 9, QFont.Bold))

            g_layout = QVBoxLayout(group)
            g_layout.setSpacing(6)

            # 불러온 값이 없으면 Config의 기본 최적값 세팅을 가져옴
            curr = self.calib_dict.get(i, Config.DEFAULT_SENSOR_CALIBRATION[i])
            ths, lss, los, hss, hos = curr["threshold"], curr["low_scale"], curr["low_offset"], curr["high_scale"], curr["high_offset"]

            self.inputs[i] = {
                "th_10": QLineEdit(str(ths[0])), "ls_10": QLineEdit(str(lss[0])), "lo_10": QLineEdit(str(los[0])), "hs_10": QLineEdit(str(hss[0])), "ho_10": QLineEdit(str(hos[0])),
                "th_25": QLineEdit(str(ths[1])), "ls_25": QLineEdit(str(lss[1])), "lo_25": QLineEdit(str(los[1])), "hs_25": QLineEdit(str(hss[1])), "ho_25": QLineEdit(str(hos[1])),
                "th_1": QLineEdit(str(ths[2])), "ls_1": QLineEdit(str(lss[2])), "lo_1": QLineEdit(str(los[2])), "hs_1": QLineEdit(str(hss[2])), "ho_1": QLineEdit(str(hos[2]))
            }

            rows = [
                ("PM10", "th_10", "ls_10", "lo_10", "hs_10", "ho_10"),
                ("PM2.5", "th_25", "ls_25", "lo_25", "hs_25", "ho_25"),
                ("PM1.0", "th_1", "ls_1", "lo_1", "hs_1", "ho_1"),
            ]

            for target_name, key_th, key_ls, key_lo, key_hs, key_ho in rows:
                row_layout = QHBoxLayout()
                lbl = QLabel(f"• {target_name}")
                lbl.setFixedWidth(50)
                lbl.setFont(QFont("Malgun Gothic", 9))
                row_layout.addWidget(lbl)

                for key, name in [(key_th, "TH:"), (key_ls, "Low S:"), (key_lo, "Low O:"), (key_hs, "High S:"), (key_ho, "High O:")]:
                    row_layout.addWidget(QLabel(name))
                    input_field = self.inputs[i][key]
                    # 🟢 [보완] 소수점 둘째 자리(-10.12 등)가 잘리지 않도록 입력칸 너비를 55로 넓힘
                    input_field.setFixedWidth(55) 
                    row_layout.addWidget(input_field)

                g_layout.addLayout(row_layout)
            scroll_layout.addWidget(group)

        scroll.setWidget(scroll_content)
        main_layout.addWidget(scroll)

        save_btn = QPushButton("적용 및 저장")
        save_btn.setFixedHeight(35)
        save_btn.setFont(QFont("Malgun Gothic", 10, QFont.Bold))
        save_btn.setStyleSheet("background-color: #28A745; color: white; border-radius: 4px;")
        save_btn.clicked.connect(self.save_calib)

        main_layout.addWidget(save_btn)

    def save_calib(self):
        try:
            new_calib = {}
            for i in range(Config.MAX_SENSORS):
                th10, th25, th1 = float(self.inputs[i]["th_10"].text()), float(self.inputs[i]["th_25"].text()), float(self.inputs[i]["th_1"].text())
                ls10, ls25, ls1 = float(self.inputs[i]["ls_10"].text()), float(self.inputs[i]["ls_25"].text()), float(self.inputs[i]["ls_1"].text())
                lo10, lo25, lo1 = float(self.inputs[i]["lo_10"].text()), float(self.inputs[i]["lo_25"].text()), float(self.inputs[i]["lo_1"].text())
                hs10, hs25, hs1 = float(self.inputs[i]["hs_10"].text()), float(self.inputs[i]["hs_25"].text()), float(self.inputs[i]["hs_1"].text())
                ho10, ho25, ho1 = float(self.inputs[i]["ho_10"].text()), float(self.inputs[i]["ho_25"].text()), float(self.inputs[i]["ho_1"].text())

                values = (th10, th25, th1, ls10, ls25, ls1, lo10, lo25, lo1, hs10, hs25, hs1, ho10, ho25, ho1)

                if not all(math.isfinite(x) for x in values):
                    raise ValueError

                if any(x < 0 for x in [ls10, ls25, ls1, hs10, hs25, hs1]):
                    QMessageBox.warning(self, "오류", "Scale 값은 0 이상이어야 합니다.")
                    return

                new_calib[i] = {
                    "threshold": (th10, th25, th1),
                    "low_scale": (ls10, ls25, ls1),
                    "low_offset": (lo10, lo25, lo1),
                    "high_scale": (hs10, hs25, hs1),
                    "high_offset": (ho10, ho25, ho1)
                }

            self.calib_dict = new_calib
            self.accept()

        except ValueError:
            QMessageBox.warning(self, "오류", "모든 입력값에는 올바른 숫자(실수)를 입력해주세요.")


# ============================================================
# Excel Export Worker
# ============================================================

class ExcelExportWorker(QThread):

    finished_signal = pyqtSignal(bool, str)

    def __init__(self, logger, date_str, parent=None):
        super().__init__(parent)
        self.logger = logger
        self.date_str = date_str

    def run(self):
        try:
            if self.isInterruptionRequested():
                return

            self.logger.db_manager.flush(timeout=5.0)
            success = self.logger.export_excel(self.date_str)

            if success:
                self.finished_signal.emit(True, f"{self.date_str} Excel 저장 완료")
            else:
                self.finished_signal.emit(False, f"{self.date_str} 저장할 데이터가 없습니다.")

        except Exception as e:
            self.finished_signal.emit(False, f"Excel 저장 오류: {e}")


# ============================================================
# Main Application
# ============================================================

class DustMonitorApp(QMainWindow):

    def __init__(self, slot_mapping):
        super().__init__()
        self.threads = []
        self.sensor_widgets = {}
        self.slot_mapping = slot_mapping
        self.db_path = os.path.join(BASE_DIR, "Data", "dust_measurement.db")

        self.export_workers = []
        self.exported_dates = set()

        self.initUI()
        self.start_monitoring()

    def initUI(self):
        self.setWindowTitle("다중 미세먼지 모니터링 시스템")
        self.setStyleSheet("QMainWindow { background-color: white; }")
        self.setMinimumSize(400, 150)

        central_widget = QWidget()
        main_layout = QVBoxLayout(central_widget)

        top_control_layout = QHBoxLayout()
        self.calib_btn = QPushButton("센서 보정 설정")
        self.calib_btn.setFont(QFont("Malgun Gothic", 9, QFont.Bold))
        self.calib_btn.setFixedHeight(30)
        self.calib_btn.clicked.connect(self.open_calibration_dialog)

        self.reload_calib_btn = QPushButton("🔄 외부 보정값 파일 새로고침")
        self.reload_calib_btn.setFont(QFont("Malgun Gothic", 9, QFont.Bold))
        self.reload_calib_btn.setFixedHeight(30)
        self.reload_calib_btn.setStyleSheet("background-color: #17A2B8; color: white; border-radius: 4px; padding: 0 10px;")
        self.reload_calib_btn.clicked.connect(self.reload_calibration_from_file)

        top_control_layout.addWidget(self.calib_btn)
        top_control_layout.addWidget(self.reload_calib_btn)
        top_control_layout.addStretch(1)

        main_layout.addLayout(top_control_layout)

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll_content = QWidget()
        scroll.setWidget(scroll_content)

        self.scroll_layout = QVBoxLayout(scroll_content)
        self.scroll_layout.setContentsMargins(15, 15, 15, 15)
        self.scroll_layout.setSpacing(15)
        self.scroll_layout.addStretch(1)

        main_layout.addWidget(scroll)
        self.setCentralWidget(central_widget)

    def open_calibration_dialog(self):
        dialog = CalibrationDialog(Config.SENSOR_CALIBRATION, self)
        if dialog.exec_() == QDialog.Accepted:
            Config.SENSOR_CALIBRATION = dialog.calib_dict
            saved = Config.save_calibration()

            if not saved:
                QMessageBox.warning(self, "오류", "보정값 파일 저장에 실패했습니다.")
                return

            apply_failed = []
            for thread in self.threads:
                idx = thread.sensor_index
                calib = Config.SENSOR_CALIBRATION.get(idx)
                if calib is not None:
                    if not thread.update_calibration(calib):
                        apply_failed.append(idx + 1)

            if apply_failed:
                QMessageBox.warning(
                    self, "경고",
                    f"보정값 파일은 저장되었지만 일부 센서에 적용하지 못했습니다.\n센서: {apply_failed}"
                )
            else:
                QMessageBox.information(self, "성공", "구간별 보정값이 저장되고 현재 센서에 즉시 반영되었습니다.")

    def reload_calibration_from_file(self):
        Config.load_calibration()
        
        apply_failed = []
        for thread in self.threads:
            idx = thread.sensor_index
            calib = Config.SENSOR_CALIBRATION.get(idx)
            if calib is not None:
                if not thread.update_calibration(calib):
                    apply_failed.append(idx + 1)
        
        if apply_failed:
            QMessageBox.warning(self, "경고", f"보정값은 읽어왔으나 일부 센서에 적용하지 못했습니다.\n센서: {apply_failed}")
        else:
            QMessageBox.information(self, "성공", "외부 파일(sensor_calibration.json)의 최신 보정값을 성공적으로 불러와 모든 센서에 적용했습니다.")

    def start_monitoring(self):
        for sensor_index, port_name in self.slot_mapping.items():
            group_box = QGroupBox(f"센서 {sensor_index + 1} 포트: {port_name}")
            group_box.setFont(QFont("Malgun Gothic", 11, QFont.Bold))

            port_layout = QHBoxLayout(group_box)
            port_layout.setSpacing(10)

            w_pm10 = DustLevelWidget("PM10 (미세먼지)", Config.PM10_LEVELS)
            w_pm25 = DustLevelWidget("PM2.5 (초미세먼지)", Config.PM25_LEVELS)
            w_pm1 = DustLevelWidget("PM1.0 (극초미세먼지)", Config.PM1_LEVELS)

            port_layout.addWidget(w_pm10)
            port_layout.addWidget(w_pm25)
            port_layout.addWidget(w_pm1)

            self.scroll_layout.insertWidget(self.scroll_layout.count() - 1, group_box)

            self.sensor_widgets[sensor_index] = {
                "pm10": w_pm10,
                "pm25": w_pm25,
                "pm1": w_pm1
            }

            thread = SerialThread(port_name, sensor_index)
            thread.data_signal.connect(self.update_data)
            thread.error_signal.connect(self.handle_error)
            thread.date_changed_signal.connect(self.handle_date_changed)

            thread.start()
            self.threads.append(thread)

        sensor_count = len(self.slot_mapping)
        self.resize(980, max(220, sensor_count * 245 + 60))

    def update_data(self, sensor_index, pm10, pm25, pm1):
        widgets = self.sensor_widgets.get(sensor_index)
        if not widgets:
            return
        widgets["pm10"].update_val(pm10)
        widgets["pm25"].update_val(pm25)
        widgets["pm1"].update_val(pm1)

    def handle_error(self, sensor_index, error_msg):
        widgets = self.sensor_widgets.get(sensor_index)
        if not widgets:
            return
        widgets["pm10"].set_error_state(error_msg)
        widgets["pm25"].set_error_state(error_msg)
        widgets["pm1"].set_error_state(error_msg)

    def handle_date_changed(self, sensor_index, previous_date):
        key = (sensor_index, previous_date)
        if key in self.exported_dates:
            return

        self.exported_dates.add(key)
        target_thread = next((t for t in self.threads if t.sensor_index == sensor_index), None)

        if target_thread is None:
            return

        worker = ExcelExportWorker(target_thread.logger, previous_date, self)
        worker.finished_signal.connect(self.handle_auto_export_finished)
        self.export_workers.append(worker)

        worker.finished.connect(lambda w=worker: self.remove_export_worker(w))
        worker.finished.connect(worker.deleteLater)
        worker.start()

    def remove_export_worker(self, worker):
        try:
            if worker in self.export_workers:
                self.export_workers.remove(worker)
        except Exception:
            pass

    def handle_auto_export_finished(self, success, message):
        if not success:
            print(f"[자동 Excel] {message}")

    def closeEvent(self, event):
        progress = QProgressDialog("프로그램을 안전하게 종료하는 중입니다...", None, 0, 0, self)
        progress.setWindowTitle("종료 중")
        progress.setWindowModality(Qt.ApplicationModal)
        progress.setCancelButton(None)
        progress.show()
        QApplication.processEvents()

        for worker in list(self.export_workers):
            if worker.isRunning():
                worker.requestInterruption()
                worker.wait(2000)

        for thread in self.threads:
            if thread.isRunning():
                thread.requestInterruption()

        for thread in self.threads:
            if thread.isRunning():
                if not thread.wait(Config.THREAD_WAIT_MS):
                    thread.logger.write_error("종료 시 센서 스레드가 제한 시간 내 종료되지 않았습니다.")

        db_manager = DatabaseManager(self.db_path)
        if not db_manager.flush(timeout=10.0):
            print("종료 전 SQLite flush에 실패했습니다.")

        today = datetime.now().strftime("%Y-%m-%d")
        for thread in self.threads:
            try:
                success = thread.logger.export_excel(today)
                if not success:
                    print(f"센서 {thread.sensor_index + 1} 당일 Excel 저장: 데이터 없음 또는 실패")
            except Exception as e:
                thread.logger.write_error(f"종료 Excel 처리 오류: {e}")

        if not db_manager.stop(timeout=5.0):
            print("SQLite Worker가 종료 제한 시간 내 종료되지 않았습니다.")

        progress.close()
        event.accept()

    def mousePressEvent(self, event):
        if event.button() == Qt.LeftButton:
            self.dragPos = event.globalPos() - self.frameGeometry().topLeft()

    def mouseMoveEvent(self, event):
        if event.buttons() == Qt.LeftButton and hasattr(self, "dragPos"):
            self.move(event.globalPos() - self.dragPos)

    def keyPressEvent(self, event):
        if event.key() == Qt.Key_Escape:
            self.close()
        else:
            super().keyPressEvent(event)


# ============================================================
# Main
# ============================================================

def main():
    app = QApplication(sys.argv)

    Config.load_calibration()

    db_path = os.path.join(BASE_DIR, "Data", "dust_measurement.db")
    SensorLogger.init_global_database(db_path)
    SensorLogger.cleanup_old_logs_global()

    dialog = PortSelectionDialog()
    if dialog.exec_() != QDialog.Accepted:
        return 0

    main_window = DustMonitorApp(dialog.selected_ports)
    main_window.show()

    return app.exec_()


if __name__ == "__main__":
    sys.exit(main())