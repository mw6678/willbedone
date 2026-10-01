import os
import glob
import sys

try:
    import pandas as pd
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, Border, Side
except ImportError:
    print("? pandas 또는 openpyxl 라이브러리가 없습니다.")
    print("터미널에서 'pip install pandas openpyxl'을 실행해주세요.")
    sys.exit(1)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
# 모니터링 프로그램이 엑셀을 저장하는 기본 경로
EXCEL_DIR = os.path.join(BASE_DIR, "Excel_Logs") 

def find_log_files():
    # 현재 폴더 및 Excel_Logs 폴더에서 Dust_log 파일 찾기
    search_paths = [BASE_DIR, EXCEL_DIR]
    candidates = []
    
    for path in search_paths:
        if os.path.exists(path):
            files = glob.glob(os.path.join(path, "Dust_log_*.xlsx"))
            candidates.extend(files)
            
    # 중복 제거 및 정렬
    candidates = sorted(list(set(candidates)))
    
    if not candidates:
        print("? 변환할 1분 단위 센서 데이터(Dust_log_...xlsx)를 찾을 수 없습니다.")
        sys.exit(1)
        
    print("\n?? 변환할 센서 데이터 파일을 선택하세요:")
    for idx, f in enumerate(candidates, 1):
        print(f"   [{idx}] {os.path.basename(f)}")
        
    while True:
        try:
            choice = input("\n번호를 입력하세요: ").strip()
            sel_idx = int(choice) - 1
            if 0 <= sel_idx < len(candidates):
                return candidates[sel_idx]
        except (ValueError, EOFError):
            pass

def convert_to_kotiti_format(input_file):
    filename = os.path.basename(input_file)
    print(f"\n? '{filename}' 파일을 읽어 1시간 평균을 계산하는 중...")
    
    try:
        # 1. 1분 단위 엑셀 파일 읽기
        df = pd.read_excel(input_file)
        
        if '측정일시' not in df.columns or 'PM2.5' not in df.columns:
            print("? 파일 형식 오류: '측정일시' 또는 'PM2.5' 열을 찾을 수 없습니다.")
            return

        # 2. 날짜/시간 데이터를 datetime으로 변환
        df['측정일시'] = pd.to_datetime(df['측정일시'])
        
        # 일(Day)과 시간(Hour) 추출
        df['일'] = df['측정일시'].dt.day
        df['시간'] = df['측정일시'].dt.hour
        
        # 3. 일/시간별 PM2.5 평균 계산 (소수점 1자리까지 반올림)
        hourly_avg = df.groupby(['일', '시간'])['PM2.5'].mean().round(1).reset_index()
        
        if hourly_avg.empty:
            print("? 계산할 데이터가 없습니다.")
            return
            
        # 4. KOTITI 양식의 새 엑셀 뼈대 만들기
        wb = Workbook()
        ws = wb.active
        ws.title = "KOTITI_Format"
        
        # 스타일 설정
        center_align = Alignment(horizontal='center', vertical='center')
        bold_font = Font(bold=True)
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                             top=Side(style='thin'), bottom=Side(style='thin'))
        
        # 헤더 1행 작성 (조회일자 병합)
        ws.merge_cells('A1:B1')
        ws['A1'] = '조회일자'
        ws['A1'].alignment = center_align
        ws['A1'].border = thin_border
        ws['B1'].border = thin_border
        
        # 헤더 2행 작성 (일, 시간, KOTITI)
        headers = ['일', '시간', 'KOTITI']
        for col_num, header_text in enumerate(headers, 1):
            cell = ws.cell(row=2, column=col_num)
            cell.value = header_text
            cell.alignment = center_align
            cell.border = thin_border
            
        # 열 너비 설정
        ws.column_dimensions['A'].width = 12
        ws.column_dimensions['B'].width = 12
        ws.column_dimensions['C'].width = 15

        # 5. 계산된 1시간 평균 데이터 채워넣기
        for row_idx, row in hourly_avg.iterrows():
            current_row = row_idx + 3
            
            # A열: "29일" 포맷
            cell_day = ws.cell(row=current_row, column=1)
            cell_day.value = f"{int(row['일'])}일"
            cell_day.alignment = center_align
            cell_day.border = thin_border
            
            # B열: "0", "1" 포맷
            cell_hour = ws.cell(row=current_row, column=2)
            cell_hour.value = int(row['시간'])
            cell_hour.alignment = center_align
            cell_hour.border = thin_border
            
            # C열: KOTITI (평균값)
            cell_val = ws.cell(row=current_row, column=3)
            cell_val.value = row['PM2.5']
            cell_val.alignment = Alignment(horizontal='right', vertical='center')
            cell_val.border = thin_border

        # 6. 파일 저장
        output_filename = filename.replace("Dust_log_", "Reference_KOTITI_")
        output_path = os.path.join(os.path.dirname(input_file), output_filename)
        
        wb.save(output_path)
        print(f"\n? 성공! KOTITI 양식으로 변환 완료.")
        print(f"저장 위치: {output_path}")

    except Exception as e:
        print(f"? 변환 중 오류가 발생했습니다: {e}")

if __name__ == "__main__":
    print("==================================================")
    print("     [1분 단위 로그 -> KOTITI 기준기 양식 변환기]     ")
    print("==================================================")
    
    target_file = find_log_files()
    convert_to_kotiti_format(target_file)