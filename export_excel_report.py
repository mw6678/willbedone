import sqlite3
import pandas as pd
import os
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

# 메인 프로그램과 동일한 DB 폴더 경로 설정[cite: 4]
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, "Data", "dust_measurement.db")

# 출력될 엑셀 파일명 (생성 시점의 시간이 파일명에 포함됨)
current_time = datetime.now().strftime("%Y%m%d_%H%M")
OUTPUT_EXCEL = os.path.join(BASE_DIR, f"시간당_센서평균_리포트_{current_time}.xlsx")

def export_to_excel():
    # DB 파일 존재 여부 확인
    db_file = DB_PATH
    if not os.path.exists(db_file):
        # 현재 폴더에 db 파일이 있는지도 체크
        if os.path.exists(os.path.join(BASE_DIR, "dust_measurement.db")):
            db_file = os.path.join(BASE_DIR, "dust_measurement.db")
        elif os.path.exists(os.path.join(BASE_DIR, "dust_measurement_2.db")):
            db_file = os.path.join(BASE_DIR, "dust_measurement_2.db") # 추가된 db 파일명 대응[cite: 5]
        else:
            print(f"❌ DB 파일을 찾을 수 없습니다: {DB_PATH}")
            return

    print("📊 DB에서 데이터를 분석하는 중입니다...")
    conn = sqlite3.connect(db_file)
    
    # 시간당 평균 쿼리 (raw_pm25 기준, 데이터 45분 이상인 유효 시간대만[cite: 4])
    query = """
    SELECT 
        strftime('%Y-%m-%d %H:00:00', measured_at) AS 시간,
        sensor_index,
        AVG(raw_pm25) AS avg_pm25
    FROM measurements
    WHERE status = 'NORMAL'
    GROUP BY sensor_index, 시간
    HAVING COUNT(*) >= 45
    ORDER BY 시간 ASC, sensor_index ASC
    """
    
    try:
        df = pd.read_sql_query(query, conn)
    except Exception as e:
        print(f"❌ 데이터 조회 실패: {e}")
        conn.close()
        return
    finally:
        conn.close()

    if df.empty:
        print("❌ 출력할 데이터가 없습니다.")
        return

    # 센서 인덱스를 보기 쉬운 이름으로 변경 (0 -> 센서 1)
    df['sensor_index'] = df['sensor_index'].apply(lambda x: f"센서 {int(x)+1}")
    
    # 데이터를 엑셀 표(피벗 테이블) 형태로 변환 (행: 시간, 열: 센서 번호)
    pivot_df = df.pivot(index='시간', columns='sensor_index', values='avg_pm25').reset_index()
    
    # 결측치(데이터가 없는 시간)는 0.0으로 채우기
    pivot_df = pivot_df.fillna(0.0)

    # 소수점 2자리까지만 반올림
    for col in pivot_df.columns:
        if col != '시간':
            pivot_df[col] = pivot_df[col].round(2)

    # 1차적으로 엑셀 파일 저장
    print("💾 엑셀 파일 서식을 적용 중입니다...")
    pivot_df.to_excel(OUTPUT_EXCEL, index=False, sheet_name="시간당 평균 데이터")
    
    # ==========================================
    # openpyxl을 이용한 엑셀 디자인 서식 적용[cite: 4]
    # ==========================================
    wb = load_workbook(OUTPUT_EXCEL)
    ws = wb.active
    
    # 스타일 정의
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, name="맑은 고딕")
    data_font = Font(name="맑은 고딕", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(left=Side(style="thin", color="D3D3D3"), right=Side(style="thin", color="D3D3D3"),
                         top=Side(style="thin", color="D3D3D3"), bottom=Side(style="thin", color="D3D3D3"))
    
    # 헤더(첫 줄) 스타일 적용
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
        
    # 데이터 셀 스타일 적용 및 열 너비 조절
    ws.column_dimensions['A'].width = 22
    for col_idx in range(2, ws.max_column + 1):
        col_letter = ws.cell(row=1, column=col_idx).column_letter
        ws.column_dimensions[col_letter].width = 15
        
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            cell.font = data_font
            cell.alignment = center_align
            cell.border = thin_border
            if isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0.00"

    # 최종 저장
    wb.save(OUTPUT_EXCEL)
    print(f"\n✅ 엑셀 추출이 완료되었습니다!")
    print(f"👉 파일 위치: {OUTPUT_EXCEL}")

if __name__ == "__main__":
    export_to_excel()